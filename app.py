import os
import sys
import socket
import json
import threading
import time
import shutil
import sqlite3
import winreg
from tkinter import Tk, Label, Entry, Button, StringVar, messagebox, ttk, filedialog, Frame, Text, Scrollbar, END, Checkbutton, IntVar, PhotoImage
import numpy as np
import pydicom
from PIL import Image
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import requests
import gc
import re
from datetime import datetime
import io
import socket as sock

try:
    import openpyxl
    from openpyxl import Workbook
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from pynetdicom import AE, evt, sop_class
from pynetdicom.presentation import AllStoragePresentationContexts
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

# ------------------------------------------------------------
# COMPRESSED DICOM CODEC SUPPORT
# ------------------------------------------------------------
CODEC_SUPPORT = {"pylibjpeg": False, "gdcm": False}
try:
    import pylibjpeg  # noqa: F401
    import pylibjpeg_libjpeg  # noqa: F401
    import pylibjpeg_openjpeg  # noqa: F401
    CODEC_SUPPORT["pylibjpeg"] = True
except Exception:
    pass
try:
    import gdcm  # noqa: F401
    CODEC_SUPPORT["gdcm"] = True
except Exception:
    pass

ACCEPTED_TRANSFER_SYNTAXES = [
    pydicom.uid.ImplicitVRLittleEndian,
    pydicom.uid.ExplicitVRLittleEndian,
    pydicom.uid.DeflatedExplicitVRLittleEndian,
    pydicom.uid.ExplicitVRBigEndian,
    pydicom.uid.JPEGBaseline8Bit,
    pydicom.uid.JPEGExtended12Bit,
    pydicom.uid.JPEGLossless,
    pydicom.uid.JPEGLosslessSV1,
    pydicom.uid.JPEGLSLossless,
    pydicom.uid.JPEGLSNearLossless,
    pydicom.uid.JPEG2000Lossless,
    pydicom.uid.JPEG2000,
    pydicom.uid.RLELossless,
]
# ------------------------------------------------------------

# ------------------------------------------------------------
# HARD‑CODED CONSTANTS
# ------------------------------------------------------------
MASTER_PASSWORD = "Sandeep@123"
DEFAULT_MASTER_USER_ID = "878604830"
DEFAULT_TELEGRAM_BOT_TOKEN = "7941135502:AAHz-KGvAAoZEhPVgfVKw3zFbkaB0_Pi5rM"
DEFAULT_WHATSAPP_API_KEY = ""
CONFIG_PASSWORD = "18040709"
DEFAULT_BATCH_WAIT_SECONDS = 6
# ------------------------------------------------------------

# --- Portable Config Directory ---
APPDATA = os.environ.get('APPDATA', os.path.expanduser('~'))
CONFIG_DIR = os.path.join(APPDATA, 'byteservices', 'RAD-XR')
CONFIG_FILE = os.path.join(CONFIG_DIR, "rad_xr_config.json")
DATABASE_DIR = os.path.join(CONFIG_DIR, "DATABASE")
DATABASE_PATH = os.path.join(DATABASE_DIR, "radxr_index.db")
FOOTER_IMAGE_PATH = os.path.join(CONFIG_DIR, "footer_image.jpg")

DEFAULT_INBOX = os.path.join(CONFIG_DIR, "Inbox")
DEFAULT_ARCHIVE = os.path.join(CONFIG_DIR, "Archive")

class DicomArchiveHandler(FileSystemEventHandler):
    def __init__(self, app_instance):
        self.app = app_instance

    def on_created(self, event):
        if not event.is_directory and event.src_path.lower().endswith('.dcm'):
            time.sleep(0.5)
            threading.Thread(target=self.app.handle_new_external_dicom, args=(event.src_path,), daemon=True).start()

    def on_moved(self, event):
        if not event.is_directory and event.dest_path.lower().endswith('.dcm'):
            time.sleep(0.5)
            threading.Thread(target=self.app.handle_new_external_dicom, args=(event.dest_path,), daemon=True).start()

class RadXrReceiverApp:
    def __init__(self, root):
        self.root = root
        self.root.title("RAD-XR")
        self.root.geometry("850x720")
        self.root.configure(bg="#1e1e24")
        
        self.autostart = '--autostart' in sys.argv
        if self.autostart:
            self.root.iconify()

        try:
            icon_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "icon.png")
            if os.path.exists(icon_path):
                icon_img = PhotoImage(file=icon_path)
                self.root.iconphoto(True, icon_img)
        except Exception:
            pass
        
        self.bg_dark = "#1e1e24"
        self.bg_card = "#2a2a35"
        self.text_light = "#f3f4f6"
        self.accent_cyan = "#06b6d4"
        self.accent_green = "#10b981"
        self.accent_red = "#ef4444"
        
        os.makedirs(CONFIG_DIR, exist_ok=True)
        os.makedirs(DATABASE_DIR, exist_ok=True)
        
        self.config = {
            "password_verified": False,
            "whatsapp_api_key": DEFAULT_WHATSAPP_API_KEY,
            "whatsapp_phone_number_id": "",
            "whatsapp_sender_phone": "",
            "institute_name": "RAD-XR IMAGING CENTER",
            "ae_title": "RAD-XR",
            "ip_address": self.get_local_ip(),
            "port": "11112",
            "receive_folder": DEFAULT_INBOX,
            "archive_folder": DEFAULT_ARCHIVE,
            "pdf_output_folder": "",
            "telegram_bot_token": DEFAULT_TELEGRAM_BOT_TOKEN,
            "footer_message": "",
            "pdf_footer_text": "",
            "pdf_footer_image": "",
            "auto_start": False,
            "bot_display_name": "RAD-XR Bot",
            "batch_wait_seconds": DEFAULT_BATCH_WAIT_SECONDS
        }
        
        self.TELEGRAM_BOT_TOKEN = None
        self.TELEGRAM_MASTER_USER_ID = DEFAULT_MASTER_USER_ID
        self.allowed_users = [DEFAULT_MASTER_USER_ID]
        self.footer_message = ""
        self.pdf_footer_text = ""
        self.pdf_footer_image = ""
        self.bot_username = ""
        self.config_unlocked = False
        
        self.server_instance = None
        self.is_listening = False
        self.bot_running = True
        self.queue_data = {}
        self.last_update_id = 0
        
        self.observer = None
        self.monitoring_active = False
        self.indexing_in_progress = False

        self.pending_batches = {}
        self.batch_lock = threading.Lock()
        self.processing_lock = threading.Lock()
        
        self.lbl_index_progress_monitor = None
        self.lbl_index_progress_config = None
        self.log_widget = None
        self.lbl_bot_name = None
        self.lbl_bot_username = None
        self.lbl_whatsapp_credits = None
        self.lbl_telegram_credits = None
        self.lock_frame = None
        self.config_pass_var = None
        self.ent_wa_phone_id = None
        self.ent_wa_sender = None
        self.auto_start_var = None
        self.ent_pdf_folder = None
        self.ent_telegram_token = None
        self.ent_ip_addr = None
        self.ent_port_num = None
        
        self.load_configuration()
        self.check_codec_support()
        self.setup_modern_styles()
        
        if self.config.get("auto_start", False):
            self._add_to_startup()
        else:
            self._remove_from_startup()
        
        self.init_db()
        self.init_telegram_users_table()
        self.init_credits_tables()
        self.init_pdf_index_table()
        self.init_admin_users_table()
        self.init_groups_table()
        self.update_bot_username()
        
        if not self.config.get("password_verified"):
            self.show_password_screen()
        else:
            self.show_main_dashboard()
            threading.Thread(target=self.index_all_existing_files, daemon=True).start()
            self.sync_archive_folder_to_dashboard()
            self.start_folder_monitor()
            self.start_telegram_bot_polling()
            if self.autostart and not self.is_listening:
                self.root.after(1000, self.toggle_server_process)

    def get_local_ip(self):
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
            s.close()
            return ip
        except Exception:
            try:
                hostname = socket.gethostname()
                for ip in socket.gethostbyname_ex(hostname)[2]:
                    if not ip.startswith("127."):
                        return ip
            except:
                pass
            return "127.0.0.1"

    # ---------- Auto‑Start ----------
    def _get_app_path(self):
        if getattr(sys, 'frozen', False):
            return sys.executable
        else:
            pythonw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
            if not os.path.exists(pythonw):
                pythonw = sys.executable
            return f'"{pythonw}" "{os.path.abspath(__file__)}"'

    def _add_to_startup(self):
        try:
            key = winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                                 r"Software\Microsoft\Windows\CurrentVersion\Run",
                                 0, winreg.KEY_SET_VALUE)
            if getattr(sys, 'frozen', False):
                cmd = f'"{self._get_app_path()}" --autostart'
            else:
                cmd = f'"{self._get_app_path()}" "{os.path.abspath(__file__)}" --autostart'
            winreg.SetValueEx(key, "RAD-XR", 0, winreg.REG_SZ, cmd)
            winreg.CloseKey(key)
            self.log_message("✅ Added to Windows Startup (with autostart)")
        except Exception as e:
            self.log_message(f"❌ Failed to add to startup: {e}")

    def _remove_from_startup(self):
        try:
            key = winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                                 r"Software\Microsoft\Windows\CurrentVersion\Run",
                                 0, winreg.KEY_SET_VALUE)
            winreg.DeleteValue(key, "RAD-XR")
            winreg.CloseKey(key)
            self.log_message("🗑️ Removed from Windows Startup")
        except FileNotFoundError:
            pass
        except Exception as e:
            self.log_message(f"❌ Failed to remove from startup: {e}")

    # ---------- Configuration ----------
    def load_configuration(self):
        if os.path.exists(CONFIG_FILE):
            try:
                with open(CONFIG_FILE, "r") as f:
                    loaded_data = json.load(f)
                    self.config.update(loaded_data)
            except Exception:
                pass
        self.config.setdefault("telegram_bot_token", DEFAULT_TELEGRAM_BOT_TOKEN)
        self.config.setdefault("footer_message", "")
        self.config.setdefault("pdf_footer_text", "")
        self.config.setdefault("pdf_footer_image", "")
        self.config.setdefault("whatsapp_api_key", DEFAULT_WHATSAPP_API_KEY)
        self.config.setdefault("whatsapp_phone_number_id", "")
        self.config.setdefault("whatsapp_sender_phone", "")
        self.config.setdefault("institute_name", "RAD-XR IMAGING CENTER")
        self.config.setdefault("ae_title", "RAD-XR")
        self.config.setdefault("ip_address", self.get_local_ip())
        self.config.setdefault("port", "11112")
        self.config.setdefault("receive_folder", DEFAULT_INBOX)
        self.config.setdefault("archive_folder", DEFAULT_ARCHIVE)
        self.config.setdefault("pdf_output_folder", "")
        self.config.setdefault("auto_start", False)
        self.config.setdefault("bot_display_name", "RAD-XR Bot")
        self.config.setdefault("batch_wait_seconds", DEFAULT_BATCH_WAIT_SECONDS)
        
        self.TELEGRAM_BOT_TOKEN = self.config["telegram_bot_token"]
        self.footer_message = self.config.get("footer_message", "")
        self.pdf_footer_text = self.config.get("pdf_footer_text", "")
        self.pdf_footer_image = self.config.get("pdf_footer_image", "")
        self.allowed_users = [DEFAULT_MASTER_USER_ID]

    def save_configuration(self):
        try:
            with open(CONFIG_FILE, "w") as f:
                json.dump(self.config, f, indent=4)
        except Exception as e:
            self.log_message(f"Error saving config: {e}")

    def setup_modern_styles(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TNotebook", background=self.bg_dark, borderwidth=0)
        style.configure("TNotebook.Tab", background=self.bg_card, foreground="#9ca3af", borderwidth=0, font=("Arial", 10, "bold"))
        style.map("TNotebook.Tab", background=[("selected", self.accent_cyan)], foreground=[("selected", self.bg_dark)])
        style.configure("Treeview", background=self.bg_card, fieldbackground=self.bg_card, foreground=self.text_light, borderwidth=0, font=("Arial", 10), rowheight=28)
        style.configure("Treeview.Heading", background="#374151", foreground=self.text_light, borderwidth=0, font=("Arial", 10, "bold"))
        style.map("Treeview", background=[("selected", "#4b5563")])
        style.configure("TFrame", background=self.bg_dark)

    def clear_screen(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    def log_message(self, msg):
        timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
        full_msg = f"[{timestamp}] {msg}"
        print(full_msg, flush=True)
        if self.log_widget:
            self.log_widget.insert(END, full_msg + "\n")
            self.log_widget.see(END)
            self.root.update_idletasks()

    # ---------- Codec support ----------
    def check_codec_support(self):
        if CODEC_SUPPORT["pylibjpeg"] or CODEC_SUPPORT["gdcm"]:
            self.log_message(
                f"✅ Compressed DICOM codec support ready "
                f"(pylibjpeg={'yes' if CODEC_SUPPORT['pylibjpeg'] else 'no'}, "
                f"gdcm={'yes' if CODEC_SUPPORT['gdcm'] else 'no'})."
            )
        else:
            self.log_message(
                "⚠️ No compressed-DICOM codec plugin found. "
                "Install: pip install pylibjpeg pylibjpeg-libjpeg pylibjpeg-openjpeg python-gdcm"
            )

    def _safe_load_pixel_array(self, ds):
        try:
            return ds.pixel_array
        except Exception as first_err:
            try:
                ds.decompress()
                return ds.pixel_array
            except Exception as second_err:
                if not (CODEC_SUPPORT["pylibjpeg"] or CODEC_SUPPORT["gdcm"]):
                    raise RuntimeError(
                        "Compressed DICOM and no decoder plugin. "
                        "Install pylibjpeg / gdcm."
                    ) from second_err
                raise RuntimeError(f"Could not decode pixel data: {second_err}") from second_err

    # ---------- Helper to normalize strings (remove extra spaces) ----------
    def _normalize_string(self, s):
        if not s:
            return ""
        s = str(s).strip()
        return re.sub(r'\s+', ' ', s)

    # ---------- Database ----------
    def init_db(self):
        os.makedirs(os.path.dirname(DATABASE_PATH), exist_ok=True)
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS dicom_index
                     (file_path TEXT PRIMARY KEY,
                      accession TEXT,
                      patient_id TEXT,
                      patient_name TEXT,
                      created_at INTEGER)''')
        c.execute('CREATE INDEX IF NOT EXISTS idx_patient_name ON dicom_index (patient_name)')
        c.execute('CREATE INDEX IF NOT EXISTS idx_patient_id ON dicom_index (patient_id)')
        conn.commit()
        conn.close()
        return DATABASE_PATH

    def init_telegram_users_table(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS telegram_users (
                        user_id TEXT PRIMARY KEY,
                        username TEXT,
                        full_name TEXT,
                        first_seen INTEGER
                     )''')
        conn.commit()
        conn.close()

    def init_admin_users_table(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS admin_users (
                        user_id TEXT PRIMARY KEY,
                        added_by TEXT,
                        added_at INTEGER
                     )''')
        conn.commit()
        conn.close()

    def init_groups_table(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS groups (
                        group_id TEXT PRIMARY KEY,
                        added_by TEXT,
                        added_at INTEGER
                     )''')
        conn.commit()
        conn.close()

    def init_credits_tables(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS whatsapp_credits (
                        id INTEGER PRIMARY KEY CHECK (id = 1),
                        remaining INTEGER DEFAULT 0
                     )''')
        c.execute("INSERT OR IGNORE INTO whatsapp_credits (id, remaining) VALUES (1, 0)")
        c.execute('''CREATE TABLE IF NOT EXISTS telegram_credits (
                        id INTEGER PRIMARY KEY CHECK (id = 1),
                        remaining INTEGER DEFAULT 0
                     )''')
        c.execute("INSERT OR IGNORE INTO telegram_credits (id, remaining) VALUES (1, 0)")
        conn.commit()
        conn.close()

    # ---------- PDF Index (cache) ----------
    def init_pdf_index_table(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''CREATE TABLE IF NOT EXISTS pdf_index (
                        lookup_key TEXT PRIMARY KEY,
                        patient_id TEXT,
                        patient_name TEXT,
                        accession TEXT,
                        pdf_path TEXT,
                        updated_at INTEGER
                     )''')
        conn.commit()
        conn.close()

    def compute_patient_key(self, patient_id, patient_name):
        patient_id = self._normalize_string(patient_id)
        patient_name = self._normalize_string(patient_name)
        if patient_id and patient_id != "N/A":
            return f"PID::{patient_id}"
        elif patient_name and patient_name != "N/A":
            return f"NAME::{patient_name.lower()}"
        else:
            return None

    def save_pdf_index(self, patient_id, patient_name, accession_no, pdf_path):
        patient_id = self._normalize_string(patient_id)
        patient_name = self._normalize_string(patient_name)
        accession_no = self._normalize_string(accession_no)
        key = self.compute_patient_key(patient_id, patient_name)
        if not key:
            return
        self.init_pdf_index_table()
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''INSERT OR REPLACE INTO pdf_index
                     (lookup_key, patient_id, patient_name, accession, pdf_path, updated_at)
                     VALUES (?, ?, ?, ?, ?, ?)''',
                  (key, patient_id, patient_name, accession_no, pdf_path, int(time.time())))
        conn.commit()
        conn.close()

    def get_saved_pdf_for_patient(self, q_id, q_name):
        self.init_pdf_index_table()
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        q_id_clean = self._normalize_string(q_id)
        q_name_clean = self._normalize_string(q_name).lower()
        name_prefix = q_name_clean[:4] if len(q_name_clean) >= 4 else q_name_clean
        c.execute("""
            SELECT pdf_path, patient_id, patient_name, accession FROM pdf_index
            WHERE patient_id = ? AND LOWER(patient_name) LIKE ?
        """, (q_id_clean, f"{name_prefix}%"))
        row = c.fetchone()
        conn.close()
        if row and row[0] and os.path.exists(row[0]):
            return row
        return None

    def clear_pdf_index_table(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("DELETE FROM pdf_index")
        conn.commit()
        conn.close()

    # ---------- Admin Users ----------
    def add_admin_user(self, user_id, added_by):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("INSERT OR REPLACE INTO admin_users (user_id, added_by, added_at) VALUES (?, ?, ?)",
                  (user_id, added_by, int(time.time())))
        conn.commit()
        conn.close()

    def remove_admin_user(self, user_id):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("DELETE FROM admin_users WHERE user_id = ?", (user_id,))
        conn.commit()
        conn.close()

    def get_admin_users(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT user_id FROM admin_users")
        rows = c.fetchall()
        conn.close()
        return [row[0] for row in rows]

    def get_all_telegram_users(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT user_id FROM telegram_users")
        results = [row[0] for row in c.fetchall()]
        conn.close()
        return results

    # ---------- Groups ----------
    def add_group(self, group_id, added_by):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("INSERT OR REPLACE INTO groups (group_id, added_by, added_at) VALUES (?, ?, ?)",
                  (group_id, added_by, int(time.time())))
        conn.commit()
        conn.close()

    def remove_group(self, group_id):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("DELETE FROM groups WHERE group_id = ?", (group_id,))
        conn.commit()
        conn.close()

    def get_all_groups(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT group_id FROM groups")
        rows = c.fetchall()
        conn.close()
        return [row[0] for row in rows]

    def get_all_groups_with_info(self):
        """Return list of (group_id, added_by, added_at) for all groups."""
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT group_id, added_by, added_at FROM groups")
        rows = c.fetchall()
        conn.close()
        return rows

    def _format_group_list(self):
        """Format group list for display."""
        groups = self.get_all_groups_with_info()
        if not groups:
            return "No groups added yet."
        lines = ["📋 *Added Groups:*", ""]
        for gid, added_by, added_at in groups:
            added_time = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(added_at)) if added_at else "Unknown"
            lines.append(f"• `{gid}`  (added by `{added_by}` on {added_time})")
        return "\n".join(lines)

    # ---------- PDF Folder & Naming ----------
    def get_pdf_folder(self):
        custom_folder = self.config.get("pdf_output_folder", "").strip()
        if custom_folder:
            try:
                os.makedirs(custom_folder, exist_ok=True)
                return custom_folder
            except Exception as e:
                self.log_message(f"⚠️ Could not create custom PDF folder '{custom_folder}': {e}. Falling back to default.")
        receive_folder = self.config.get("receive_folder", DEFAULT_INBOX)
        parent_dir = os.path.dirname(os.path.normpath(receive_folder))
        pdf_dir = os.path.join(parent_dir, "PDF Reports")
        os.makedirs(pdf_dir, exist_ok=True)
        return pdf_dir

    def build_pdf_filename(self, patient_name, study_date):
        patient_name = self._normalize_string(patient_name)
        if not patient_name:
            patient_name = "Unknown"
        if study_date and len(study_date) == 8 and study_date.isdigit():
            formatted_date = f"{study_date[6:8]}{study_date[4:6]}{study_date[0:4]}"
        else:
            today = time.strftime("%Y%m%d")
            formatted_date = f"{today[6:8]}{today[4:6]}{today[0:4]}"
        safe_name = "".join(c for c in patient_name if c.isalnum() or c in " -_")
        return f"{safe_name} {formatted_date} medical report.pdf"

    # ---------- Credits ----------
    def get_whatsapp_credits(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT remaining FROM whatsapp_credits WHERE id = 1")
        row = c.fetchone()
        conn.close()
        return row[0] if row else 0

    def add_whatsapp_credits(self, amount):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("UPDATE whatsapp_credits SET remaining = remaining + ? WHERE id = 1", (amount,))
        conn.commit()
        new_total = c.execute("SELECT remaining FROM whatsapp_credits WHERE id = 1").fetchone()[0]
        conn.close()
        return new_total

    def decrement_whatsapp_credits(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT remaining FROM whatsapp_credits WHERE id = 1")
        row = c.fetchone()
        if row and row[0] > 0:
            c.execute("UPDATE whatsapp_credits SET remaining = remaining - 1 WHERE id = 1")
            conn.commit()
            conn.close()
            return True
        conn.close()
        return False

    def get_telegram_credits(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT remaining FROM telegram_credits WHERE id = 1")
        row = c.fetchone()
        conn.close()
        return row[0] if row else 0

    def add_telegram_credits(self, amount):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("UPDATE telegram_credits SET remaining = remaining + ? WHERE id = 1", (amount,))
        conn.commit()
        new_total = c.execute("SELECT remaining FROM telegram_credits WHERE id = 1").fetchone()[0]
        conn.close()
        return new_total

    def decrement_telegram_credits(self):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("SELECT remaining FROM telegram_credits WHERE id = 1")
        row = c.fetchone()
        if row and row[0] > 0:
            c.execute("UPDATE telegram_credits SET remaining = remaining - 1 WHERE id = 1")
            conn.commit()
            conn.close()
            return True
        conn.close()
        return False

    # ---------- Telegram Users (regular) ----------
    def save_telegram_user(self, user_id, username, full_name):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute('''INSERT OR REPLACE INTO telegram_users (user_id, username, full_name, first_seen)
                     VALUES (?, ?, ?, COALESCE((SELECT first_seen FROM telegram_users WHERE user_id=?), ?))''',
                  (user_id, username, full_name, user_id, int(time.time())))
        conn.commit()
        conn.close()

    def delete_telegram_user(self, user_id):
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("DELETE FROM telegram_users WHERE user_id = ?", (user_id,))
        conn.commit()
        conn.close()

    # ---------- Bot info ----------
    def update_bot_username(self):
        if not self.TELEGRAM_BOT_TOKEN:
            return
        try:
            url = f"https://api.telegram.org/bot{self.TELEGRAM_BOT_TOKEN}/getMe"
            resp = requests.get(url, timeout=10)
            if resp.status_code == 200:
                data = resp.json()
                if data.get("ok"):
                    self.bot_username = data["result"].get("username", "")
                    self.log_message(f"Bot username: @{self.bot_username}")
                    self.config["bot_display_name"] = data["result"].get("first_name", "RAD-XR Bot")
                    self.save_configuration()
                    self.refresh_bot_info_gui()
        except Exception as e:
            self.log_message(f"Failed to fetch bot info: {e}")

    def refresh_bot_info_gui(self):
        if self.lbl_bot_name:
            self.lbl_bot_name.config(text=f"Name: {self.config.get('bot_display_name', 'RAD-XR Bot')}")
        if self.lbl_bot_username:
            bot_uname = f"@{self.bot_username}" if self.bot_username else "Not available"
            self.lbl_bot_username.config(text=f"Username: {bot_uname}")

    def refresh_credits_gui(self):
        if self.lbl_whatsapp_credits:
            self.lbl_whatsapp_credits.config(text=f"💰 WhatsApp: {self.get_whatsapp_credits()}")
        if self.lbl_telegram_credits:
            self.lbl_telegram_credits.config(text=f"🤖 Telegram: {self.get_telegram_credits()}")

    # ---------- Indexing (now indexes Archive folder) ----------
    def index_all_existing_files(self, reindex=False):
        if self.indexing_in_progress:
            self.log_message("Indexing already in progress. Skipping.")
            return
        self.indexing_in_progress = True
        self.log_message("Starting indexing of Archive folder...")
        archive_dir = self.config["archive_folder"]
        if not os.path.exists(archive_dir):
            self.log_message(f"Archive folder {archive_dir} does not exist.")
            self.indexing_in_progress = False
            return

        self.init_db()
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()

        if reindex:
            self.log_message("Re-indexing: dropping existing table...")
            c.execute("DROP TABLE IF EXISTS dicom_index")
            self.init_db()
            conn = sqlite3.connect(DATABASE_PATH)
            c = conn.cursor()

        all_files = [f for f in os.listdir(archive_dir) if f.lower().endswith(".dcm")]
        total = len(all_files)
        processed = 0

        self.root.after(0, self.update_index_progress, processed, total)

        for file in all_files:
            full_path = os.path.join(archive_dir, file)
            try:
                if not reindex:
                    c.execute("SELECT file_path FROM dicom_index WHERE file_path = ?", (full_path,))
                    if c.fetchone():
                        processed += 1
                        self.root.after(0, self.update_index_progress, processed, total)
                        continue
                ds = pydicom.dcmread(full_path, stop_before_pixels=True)
                accession = self._normalize_string(ds.get("AccessionNumber", "UNKNOWN"))
                patient_id = self._normalize_string(ds.get("PatientID", "N/A"))
                patient_name = self._normalize_string(ds.get("PatientName", "N/A"))
                c.execute("INSERT OR IGNORE INTO dicom_index (file_path, accession, patient_id, patient_name, created_at) VALUES (?,?,?,?,?)",
                          (full_path, accession, patient_id, patient_name, int(os.path.getctime(full_path))))
            except Exception as e:
                self.log_message(f"❌ Indexing failed for {full_path}: {e}")
            processed += 1
            self.root.after(0, self.update_index_progress, processed, total)

        conn.commit()
        conn.close()
        self.indexing_in_progress = False
        self.root.after(0, self.update_index_progress, total, total, done=True)
        self.log_message(f"✅ Indexing complete: {processed} files indexed.")

    def update_index_progress(self, current, total, done=False):
        if done:
            text = f"✅ Indexing complete: {total} files indexed."
        else:
            text = f"⏳ Indexing: {current} / {total} files processed..."
        if self.lbl_index_progress_monitor:
            self.lbl_index_progress_monitor.config(text=text)
        if self.lbl_index_progress_config:
            self.lbl_index_progress_config.config(text=text)

    def index_dicom_file(self, dcm_path, patient_id, patient_name, accession_no):
        if not os.path.exists(dcm_path):
            return
        patient_id = self._normalize_string(patient_id)
        patient_name = self._normalize_string(patient_name)
        accession_no = self._normalize_string(accession_no)
        self.init_db()
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("INSERT OR IGNORE INTO dicom_index (file_path, accession, patient_id, patient_name, created_at) VALUES (?,?,?,?,?)",
                  (dcm_path, accession_no, patient_id, patient_name, int(time.time())))
        conn.commit()
        conn.close()

    def remove_from_index(self, file_path):
        self.init_db()
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        c.execute("DELETE FROM dicom_index WHERE file_path = ?", (file_path,))
        conn.commit()
        conn.close()

    def get_patient_files_from_db(self, q_id, q_name):
        self.init_db()
        conn = sqlite3.connect(DATABASE_PATH)
        c = conn.cursor()
        q_id_clean = self._normalize_string(q_id)
        q_name_clean = self._normalize_string(q_name).lower()
        name_prefix = q_name_clean[:4] if len(q_name_clean) >= 4 else q_name_clean
        c.execute("""
            SELECT file_path, patient_id, patient_name, accession
            FROM dicom_index
            WHERE patient_id = ? AND LOWER(patient_name) LIKE ?
        """, (q_id_clean, f"{name_prefix}%"))
        results = c.fetchall()
        conn.close()
        return results

    # ---------- Folder Monitoring (for Archive, external) ----------
    def start_folder_monitor(self):
        archive_dir = self.config["archive_folder"]
        try:
            os.makedirs(archive_dir, exist_ok=True)
        except Exception as e:
            self.log_message(f"❌ Failed to create archive folder: {e}")
            messagebox.showerror("Folder Error", f"Cannot create archive folder:\n{archive_dir}\n\nError: {e}")
            return
        if self.observer and self.monitoring_active:
            self.stop_folder_monitor()
        self.observer = Observer()
        event_handler = DicomArchiveHandler(self)
        self.observer.schedule(event_handler, path=archive_dir, recursive=False)
        self.observer.start()
        self.monitoring_active = True
        self.log_message(f"📁 Watching Archive folder (external): {archive_dir}")

    def stop_folder_monitor(self):
        if self.observer:
            self.observer.stop()
            self.observer.join()
            self.monitoring_active = False
            self.log_message("📁 Folder monitoring stopped.")

    def handle_new_external_dicom(self, file_path):
        try:
            if not os.path.exists(file_path):
                return
            self.queue_file_for_patient_batch(file_path, is_manual_import=False)
        except Exception as e:
            self.log_message(f"❌ Error processing external DICOM {file_path}: {e}")

    # ---------- GUI ----------
    def show_password_screen(self):
        self.clear_screen()
        main_card = Frame(self.root, bg=self.bg_card, bd=0)
        main_card.place(relx=0.5, rely=0.5, anchor="center", width=420, height=350)
        Label(main_card, text="RAD-XR SYSTEM NODE", font=("Arial", 18, "bold"), fg=self.accent_cyan, bg=self.bg_card).pack(pady=(40, 5))
        Label(main_card, text="Enterprise Activation Gateway", font=("Arial", 10), fg="#9ca3af", bg=self.bg_card).pack(pady=(0, 25))
        Label(main_card, text="Node Master Password:", font=("Arial", 10, "bold"), fg=self.text_light, bg=self.bg_card).pack(anchor="w", padx=45, pady=5)
        self.pass_var = StringVar()
        entry_pass = Entry(main_card, textvariable=self.pass_var, show="*", font=("Arial", 12), width=28, justify="center", bg=self.bg_dark, fg=self.text_light, bd=1, insertbackground="white")
        entry_pass.pack(pady=5)
        entry_pass.focus()
        btn_verify = Button(main_card, text="Verify Node", font=("Arial", 11, "bold"), bg=self.accent_cyan, fg=self.bg_dark, width=18, activebackground="#0891b2", bd=0, cursor="hand2", command=self.verify_master_password)
        btn_verify.pack(pady=35)

    def verify_master_password(self):
        if self.pass_var.get() == MASTER_PASSWORD:
            self.config["password_verified"] = True
            self.save_configuration()
            messagebox.showinfo("Success", "RAD-XR Node Authenticated!")
            self.show_main_dashboard()
            threading.Thread(target=self.index_all_existing_files, daemon=True).start()
            self.sync_archive_folder_to_dashboard()
            self.start_folder_monitor()
            self.start_telegram_bot_polling()
            if self.autostart and not self.is_listening:
                self.root.after(1000, self.toggle_server_process)
        else:
            messagebox.showerror("Error", "Invalid Security Master Password!")

    def show_main_dashboard(self):
        self.clear_screen()
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=15, pady=15)
        frame_receiver = Frame(self.notebook, bg=self.bg_dark)
        frame_settings = Frame(self.notebook, bg=self.bg_dark)
        self.notebook.add(frame_receiver, text="  Live Monitor  ")
        self.notebook.add(frame_settings, text="  Config Control  ")

        self.build_live_monitor_tab(frame_receiver)
        self.build_config_tab(frame_settings)

        self.notebook.bind("<<NotebookTabChanged>>", self.on_tab_change)
        self.refresh_bot_info_gui()
        self.refresh_credits_gui()

    def on_tab_change(self, event=None):
        if not self.notebook:
            return
        selected = self.notebook.select()
        tab_text = self.notebook.tab(selected, "text")
        if tab_text == "  Config Control  " and self.lock_frame:
            self.config_unlocked = False
            self.lock_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
            if self.config_pass_var:
                self.config_pass_var.set("")
            for child in self.lock_frame.winfo_children():
                if isinstance(child, Frame):
                    for sub in child.winfo_children():
                        if isinstance(sub, Entry):
                            sub.focus()
                            break

    def unlock_config_tab(self):
        if self.config_pass_var.get() == CONFIG_PASSWORD:
            self.config_unlocked = True
            if self.lock_frame:
                self.lock_frame.place_forget()
            messagebox.showinfo("Success", "Config Control unlocked!")
        else:
            messagebox.showerror("Error", "Incorrect password!")
            self.config_pass_var.set("")
            for child in self.lock_frame.winfo_children():
                if isinstance(child, Frame):
                    for sub in child.winfo_children():
                        if isinstance(sub, Entry):
                            sub.focus()
                            break

    def build_live_monitor_tab(self, parent):
        top_ctrl_bar = Frame(parent, bg=self.bg_dark)
        top_ctrl_bar.pack(fill="x", pady=15, padx=15)
        Label(top_ctrl_bar, text="RAD-XR", font=("Arial", 14, "bold"), fg=self.text_light, bg=self.bg_dark).pack(side="left")
        self.status_var = StringVar(value="● Stopped")
        self.lbl_status_indicator = Label(top_ctrl_bar, textvariable=self.status_var, font=("Arial", 11, "bold"), fg=self.accent_red, bg=self.bg_dark)
        self.lbl_status_indicator.pack(side="left", padx=20)

        btn_manual_upload = Button(top_ctrl_bar, text="+ Import DICOM File", font=("Arial", 9, "bold"), bg=self.accent_cyan, fg=self.bg_dark, bd=0, padx=10, pady=5, cursor="hand2", command=self.manual_file_upload_trigger)
        btn_manual_upload.pack(side="right", padx=5)

        btn_resend_failed = Button(top_ctrl_bar, text="🔄 Resend Failed", font=("Arial", 9, "bold"), bg="#f59e0b", fg=self.bg_dark, bd=0, padx=10, pady=5, cursor="hand2", command=self.resend_failed_images)
        btn_resend_failed.pack(side="right", padx=5)

        btn_clear_exams = Button(top_ctrl_bar, text="🧹 Clear Exams", font=("Arial", 9, "bold"), bg=self.accent_red, fg=self.text_light, bd=0, padx=10, pady=5, cursor="hand2", command=self.clear_exams_action)
        btn_clear_exams.pack(side="right", padx=5)

        self.btn_toggle_server = Button(top_ctrl_bar, text="Start Server", bg=self.accent_green, fg=self.bg_dark, font=("Arial", 9, "bold"), bd=0, padx=10, pady=5, width=12, cursor="hand2", command=self.toggle_server_process)
        self.btn_toggle_server.pack(side="right", padx=5)

        content_splitter = Frame(parent, bg=self.bg_dark)
        content_splitter.pack(fill="both", expand=True, padx=15, pady=5)

        net_card = Frame(content_splitter, bg=self.bg_card, width=220)
        net_card.pack(side="left", fill="y", padx=(0, 10))
        net_card.pack_propagate(False)
        Label(net_card, text="NETWORK CONFIG", font=("Arial", 10, "bold"), fg=self.accent_cyan, bg=self.bg_card).pack(pady=15)

        def add_stat_lbl(parent, name, val_attr):
            Label(parent, text=name, font=("Arial", 8, "bold"), fg="#9ca3af", bg=self.bg_card).pack(anchor="w", padx=15, pady=(5, 0))
            lbl_v = Label(parent, text=self.config.get(val_attr, "N/A"), font=("Arial", 9), fg=self.text_light, bg=self.bg_card, wraplength=190, justify="left")
            lbl_v.pack(anchor="w", padx=15, pady=(0, 5))
            return lbl_v

        self.lbl_inst = add_stat_lbl(net_card, "INSTITUTE NAME", "institute_name")
        self.lbl_ae = add_stat_lbl(net_card, "AE TITLE", "ae_title")
        self.lbl_ip = add_stat_lbl(net_card, "IP ADDRESS", "ip_address")
        self.lbl_port = add_stat_lbl(net_card, "PORT NUMBER", "port")

        Label(net_card, text="💳 CREDITS", font=("Arial", 8, "bold"), fg=self.accent_cyan, bg=self.bg_card).pack(anchor="w", padx=15, pady=(10, 0))
        self.lbl_whatsapp_credits = Label(net_card, text=f"💰 WhatsApp: {self.get_whatsapp_credits()}", font=("Arial", 9), fg=self.text_light, bg=self.bg_card, wraplength=190, justify="left")
        self.lbl_whatsapp_credits.pack(anchor="w", padx=15, pady=(0, 2))
        self.lbl_telegram_credits = Label(net_card, text=f"🤖 Telegram: {self.get_telegram_credits()}", font=("Arial", 9), fg=self.text_light, bg=self.bg_card, wraplength=190, justify="left")
        self.lbl_telegram_credits.pack(anchor="w", padx=15, pady=(0, 10))

        Label(net_card, text="🤖 TELEGRAM BOT", font=("Arial", 8, "bold"), fg=self.accent_cyan, bg=self.bg_card).pack(anchor="w", padx=15, pady=(10, 0))
        bot_display = self.config.get("bot_display_name", "RAD-XR Bot")
        self.lbl_bot_name = Label(net_card, text=f"Name: {bot_display}", font=("Arial", 9), fg=self.text_light, bg=self.bg_card, wraplength=190, justify="left")
        self.lbl_bot_name.pack(anchor="w", padx=15, pady=(0, 2))
        bot_uname = f"@{self.bot_username}" if self.bot_username else "Not available"
        self.lbl_bot_username = Label(net_card, text=f"Username: {bot_uname}", font=("Arial", 9), fg="#9ca3af", bg=self.bg_card, wraplength=190, justify="left")
        self.lbl_bot_username.pack(anchor="w", padx=15, pady=(0, 5))
        Label(net_card, text="Search on Telegram: @...", font=("Arial", 7), fg="#6b7280", bg=self.bg_card).pack(anchor="w", padx=15, pady=(0, 5))

        Label(net_card, text="📁 FOLDER WATCH", font=("Arial", 8, "bold"), fg=self.accent_green if self.monitoring_active else self.accent_red, bg=self.bg_card).pack(anchor="w", padx=15, pady=(10,0))
        Label(net_card, text="ACTIVE" if self.monitoring_active else "INACTIVE", font=("Arial", 9, "bold"), fg=self.accent_green if self.monitoring_active else self.accent_red, bg=self.bg_card).pack(anchor="w", padx=15, pady=(0, 10))
        Button(net_card, text="Refresh Dashboard", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0, width=20, cursor="hand2", command=self.refresh_dashboard).pack(side="bottom", pady=20)

        right_panel = Frame(content_splitter, bg=self.bg_card)
        right_panel.pack(side="right", fill="both", expand=True)
        queue_container = Frame(right_panel, bg=self.bg_card)
        queue_container.pack(fill="both", expand=True)

        cols = ("id", "name", "accession", "file", "status")
        self.tree = ttk.Treeview(queue_container, columns=cols, show="headings")
        self.tree.heading("id", text="Patient ID")
        self.tree.heading("name", text="Patient Name")
        self.tree.heading("accession", text="Accession No.")
        self.tree.heading("file", text="File Name")
        self.tree.heading("status", text="Status")
        self.tree.column("id", width=90, anchor="center")
        self.tree.column("name", width=140, anchor="w")
        self.tree.column("accession", width=100, anchor="center")
        self.tree.column("file", width=120, anchor="w")
        self.tree.column("status", width=160, anchor="center")
        self.tree.pack(fill="both", expand=True, side="left")
        self.tree.bind("<Double-1>", self.on_grid_row_double_click_resend)
        scrollbar = ttk.Scrollbar(queue_container, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")

        log_frame = Frame(right_panel, bg=self.bg_dark, height=120)
        log_frame.pack(fill="x", pady=(5,0))
        log_frame.pack_propagate(False)
        log_header = Frame(log_frame, bg=self.bg_dark)
        log_header.pack(fill="x")
        Label(log_header, text="📋 Console Log", font=("Arial", 9, "bold"), fg=self.accent_cyan, bg=self.bg_dark).pack(side="left")
        Button(log_header, text="📋 Copy Log", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0, padx=5, pady=2, cursor="hand2", command=self.copy_log).pack(side="right", padx=2)
        Button(log_header, text="🗑️ Clear Log", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0, padx=5, pady=2, cursor="hand2", command=self.clear_log).pack(side="right", padx=2)
        log_text_frame = Frame(log_frame, bg=self.bg_dark)
        log_text_frame.pack(fill="both", expand=True)
        self.log_widget = Text(log_text_frame, bg=self.bg_dark, fg=self.text_light, font=("Courier", 8), wrap="word", height=5)
        self.log_widget.pack(side="left", fill="both", expand=True)
        log_scroll = Scrollbar(log_text_frame, command=self.log_widget.yview)
        log_scroll.pack(side="right", fill="y")
        self.log_widget.config(yscrollcommand=log_scroll.set)

        progress_frame = Frame(parent, bg=self.bg_dark)
        progress_frame.pack(fill="x", padx=15, pady=5)
        self.lbl_index_progress_monitor = Label(progress_frame, text="✅ Indexing ready.", font=("Arial", 9), fg=self.accent_green, bg=self.bg_dark)
        self.lbl_index_progress_monitor.pack(side="left")

        btn_clear_sent = Button(progress_frame, text="🧹 Clear Sent Files", font=("Arial", 9, "bold"), bg="#3b82f6", fg=self.text_light, bd=0, padx=10, pady=2, cursor="hand2", command=self.clear_sent_files_from_grid)
        btn_clear_sent.pack(side="left", padx=10)

        Label(progress_frame, text="Made with ❤️ by Sandeep", font=("Arial", 9, "bold", "italic"), fg="#6b7280", bg=self.bg_dark).pack(side="right")

    def refresh_dashboard(self):
        self.refresh_credits_gui()
        self.sync_archive_folder_to_dashboard()
        self.log_message("🔄 Dashboard refreshed (credits updated).")

    # --- Clear Sent Files (fully sent) ---
    def clear_sent_files_from_grid(self):
        to_remove = []
        for item in self.tree.get_children():
            values = self.tree.item(item, 'values')
            if not values:
                continue
            status = values[4]
            if '❌' not in status:
                to_remove.append(item)
        if not to_remove:
            messagebox.showinfo("Clear Sent", "No fully sent files to clear.")
            return
        removed_count = 0
        for item in to_remove:
            file_path_to_remove = None
            for fpath, rid in list(self.queue_data.items()):
                if rid == item:
                    file_path_to_remove = fpath
                    break
            if file_path_to_remove:
                del self.queue_data[file_path_to_remove]
            self.tree.delete(item)
            removed_count += 1
        self.log_message(f"🧹 Cleared {removed_count} fully sent file(s) from the grid.")
        messagebox.showinfo("Clear Sent", f"Removed {removed_count} sent file(s).")

    # ---------- Status parsing helper ----------
    def _parse_status(self, status_str):
        """
        Parse status string like "Telegram ✅, WhatsApp ❌" or "Telegram ✅"
        Returns dict: {'telegram': True/False/None, 'whatsapp': True/False/None}
        True = success, False = failure, None = not applicable
        """
        result = {'telegram': None, 'whatsapp': None}
        if not status_str:
            return result
        # Split by comma
        parts = [p.strip() for p in status_str.split(',')]
        for part in parts:
            if 'Telegram' in part:
                if '✅' in part:
                    result['telegram'] = True
                elif '❌' in part:
                    result['telegram'] = False
                else:
                    result['telegram'] = None
            elif 'WhatsApp' in part:
                if '✅' in part:
                    result['whatsapp'] = True
                elif '❌' in part:
                    result['whatsapp'] = False
                elif '⏭️' in part:
                    result['whatsapp'] = True  # treat skip as success
                else:
                    result['whatsapp'] = None
        return result

    def _build_status_string(self, tg_ok, wa_ok, wa_skip=False):
        """Build status string from platform booleans."""
        parts = []
        if self.TELEGRAM_BOT_TOKEN:
            parts.append(f"Telegram {'✅' if tg_ok else '❌'}")
        if self.config.get("whatsapp_api_key") and self.config.get("whatsapp_phone_number_id"):
            if wa_skip:
                parts.append("WhatsApp ⏭️")
            else:
                parts.append(f"WhatsApp {'✅' if wa_ok else '❌'}")
        if not parts:
            parts.append("No platforms configured")
        return ", ".join(parts)

    # ---------- Resend Failed (with batching) ----------
    def resend_failed_images(self):
        """Gather failed items, group by patient, and resend each group as a batch."""
        failed_items = []
        for item in self.tree.get_children():
            values = self.tree.item(item, 'values')
            if values and "❌" in values[4]:
                failed_items.append(item)

        if not failed_items:
            messagebox.showinfo("Resend", "No failed images to resend.")
            return

        count = len(failed_items)
        if not messagebox.askyesno("Resend Failed", f"Found {count} failed image(s). Resend all?"):
            return

        # Group by patient key
        groups = {}  # key -> list of (item_id, file_path, patient_id, patient_name, accession, study_date)
        for item in failed_items:
            row_values = self.tree.item(item, 'values')
            if not row_values or len(row_values) < 5:
                continue
            p_id = row_values[0]
            p_name = row_values[1]
            accession = row_values[2]
            file_name = row_values[3]
            # Find file path
            file_path = None
            for fpath, rid in self.queue_data.items():
                if rid == item:
                    file_path = fpath
                    break
            if not file_path:
                # Try inbox or archive
                inbox_file = os.path.join(self.config["receive_folder"], file_name)
                if os.path.exists(inbox_file):
                    file_path = inbox_file
                else:
                    archive_file = os.path.join(self.config["archive_folder"], file_name)
                    if os.path.exists(archive_file):
                        file_path = archive_file
            if not file_path or not os.path.exists(file_path):
                self.log_message(f"❌ File not found: {file_name}")
                continue
            # Read study date from DICOM
            study_date = ""
            try:
                ds = pydicom.dcmread(file_path, stop_before_pixels=True)
                study_date = self._normalize_string(ds.get("StudyDate", ""))
            except:
                pass
            key = self.compute_patient_key(p_id, p_name)
            if not key:
                key = f"UNKNOWN::{time.time_ns()}"
            if key not in groups:
                groups[key] = []
            groups[key].append({
                'item_id': item,
                'file_path': file_path,
                'patient_id': p_id,
                'patient_name': p_name,
                'accession': accession,
                'study_date': study_date,
                'file_name': file_name
            })

        # Process each group
        for key, group_items in groups.items():
            threading.Thread(target=self.resend_batch_for_patient, args=(group_items,), daemon=True).start()

        messagebox.showinfo("Resend", f"Started resending {len(groups)} patient group(s). Check console.")

    def resend_batch_for_patient(self, group_items):
        """
        Resend a group of failed files for the same patient as one combined PDF.
        group_items: list of dicts with item_id, file_path, patient_id, patient_name, accession, study_date
        """
        with self.processing_lock:
            try:
                # Prepare data
                file_paths = [item['file_path'] for item in group_items]
                patient_id = group_items[0]['patient_id']
                patient_name = group_items[0]['patient_name']
                accession_no = group_items[0]['accession']
                study_date = group_items[0]['study_date']

                # Determine which platforms are enabled
                tg_enabled = bool(self.TELEGRAM_BOT_TOKEN)
                wa_enabled = bool(self.config.get("whatsapp_api_key") and self.config.get("whatsapp_phone_number_id"))

                # Parse current status for each file to know which platforms failed
                # We'll aggregate: for each platform, if any file has a failure, we need to retry.
                tg_need_retry = False
                wa_need_retry = False
                for item in group_items:
                    status_str = self.tree.item(item['item_id'], 'values')[4]
                    parsed = self._parse_status(status_str)
                    if parsed.get('telegram') is False:
                        tg_need_retry = True
                    if parsed.get('whatsapp') is False:
                        wa_need_retry = True

                # Generate combined PDF
                pdf_folder = self.get_pdf_folder()
                pdf_path = os.path.join(pdf_folder, self.build_pdf_filename(patient_name, study_date))
                self.generate_pdf_report_from_dicom(file_paths, pdf_path)

                # Update status to "Retrying..."
                for item in group_items:
                    self.root.after(0, lambda i=item['item_id'], p=patient_id, n=patient_name, a=accession_no:
                                    self.tree.item(i, values=(p, n, a, item['file_name'], "⚡ Resending...")))

                # Retry Telegram if needed
                tg_ok = True
                if tg_enabled and tg_need_retry:
                    if self.decrement_telegram_credits():
                        tg_ok = self.send_to_all_telegram(pdf_path, patient_id, patient_name, accession_no)
                        if not tg_ok:
                            self.add_telegram_credits(1)
                            self.log_message("⚠️ Telegram resend failed, credit refunded.")
                    else:
                        tg_ok = False
                        self.log_message("⚠️ No Telegram credits for resend.")
                elif not tg_enabled:
                    tg_ok = True  # not applicable

                # Retry WhatsApp if needed (and if accession not missing)
                wa_ok = True
                wa_skip = False
                accession_missing = not accession_no or accession_no in ("N/A", "UNKNOWN")
                if wa_enabled and wa_need_retry and not accession_missing:
                    if self.decrement_whatsapp_credits():
                        wa_ok = self.dispatch_to_whatsapp_business(pdf_path, patient_id, patient_name, accession_no)
                        if not wa_ok:
                            self.add_whatsapp_credits(1)
                            self.log_message("⚠️ WhatsApp resend failed, credit refunded.")
                    else:
                        wa_ok = False
                        self.log_message("⚠️ No WhatsApp credits for resend.")
                elif wa_enabled and accession_missing:
                    wa_skip = True
                    self.log_message(f"ℹ️ WhatsApp skipped: Accession Number missing for {patient_name} ({patient_id})")
                elif not wa_enabled:
                    wa_ok = True  # not applicable

                # Build final status string
                status_str = self._build_status_string(tg_ok, wa_ok, wa_skip)

                # Update all items in group with the new status
                for item in group_items:
                    self.root.after(0, lambda i=item['item_id'], p=patient_id, n=patient_name, a=accession_no, s=status_str:
                                    self.tree.item(i, values=(p, n, a, item['file_name'], s)))

                # Check if all ok (or skipped)
                all_ok = (not tg_enabled or tg_ok) and (not wa_enabled or wa_ok or wa_skip)
                if all_ok:
                    # Save PDF index
                    self.save_pdf_index(patient_id, patient_name, accession_no, pdf_path)
                    self.log_message(f"💾 PDF saved and indexed for resend: {os.path.basename(pdf_path)}")
                    # Delete all DICOM files in the group
                    for item in group_items:
                        try:
                            if os.path.exists(item['file_path']):
                                os.remove(item['file_path'])
                                self.remove_from_index(item['file_path'])
                                self.log_message(f"🗑️ Deleted DICOM after resend: {os.path.basename(item['file_path'])}")
                        except Exception as e:
                            self.log_message(f"⚠️ Could not delete {item['file_path']}: {e}")
                else:
                    self.log_message(f"⚠️ Resend partial success for {patient_name} ({patient_id}): {status_str}")

            except Exception as e:
                self.log_message(f"❌ Resend batch error: {e}")
                # Update items to show error
                for item in group_items:
                    self.root.after(0, lambda i=item['item_id'], p=patient_id, n=patient_name, a=accession_no:
                                    self.tree.item(i, values=(p, n, a, item['file_name'], f"Error: {str(e)[:30]}")))

    # ---------- Resend single file (kept for double‑click) ----------
    def resend_single_file(self, file_path, item_id):
        with self.processing_lock:
            row_values = self.tree.item(item_id, 'values')
            status_str = row_values[4]
            tg_enabled = bool(self.TELEGRAM_BOT_TOKEN)
            wa_enabled = bool(self.config.get("whatsapp_api_key") and self.config.get("whatsapp_phone_number_id"))
            parsed = self._parse_status(status_str)
            tg_ok = parsed.get('telegram') if parsed.get('telegram') is not None else True
            wa_ok = parsed.get('whatsapp') if parsed.get('whatsapp') is not None else True
            wa_skip = False  # will be set if accession missing

            try:
                ds = pydicom.dcmread(file_path)
                patient_id = self._normalize_string(ds.get("PatientID", "N/A"))
                patient_name = self._normalize_string(ds.get("PatientName", "N/A"))
                accession_no = self._normalize_string(ds.get("AccessionNumber", "UNKNOWN"))
                study_date = self._normalize_string(ds.get("StudyDate", ""))

                pdf_folder = self.get_pdf_folder()
                pdf_path = os.path.join(pdf_folder, self.build_pdf_filename(patient_name, study_date))
                self.generate_pdf_report_from_dicom([file_path], pdf_path)

                new_tg_ok = tg_ok
                new_wa_ok = wa_ok
                new_wa_skip = False

                if not tg_ok and tg_enabled:
                    if self.decrement_telegram_credits():
                        if self.send_to_all_telegram(pdf_path, patient_id, patient_name, accession_no):
                            new_tg_ok = True
                        else:
                            self.add_telegram_credits(1)
                    else:
                        self.log_message("⚠️ No Telegram credits for resend.")

                accession_missing = not accession_no or accession_no in ("N/A", "UNKNOWN")
                if accession_missing:
                    new_wa_skip = True
                    self.log_message(f"ℹ️ WhatsApp skipped: Accession Number missing for {patient_name} ({patient_id})")
                elif not wa_ok and wa_enabled:
                    if self.decrement_whatsapp_credits():
                        if self.dispatch_to_whatsapp_business(pdf_path, patient_id, patient_name, accession_no):
                            new_wa_ok = True
                        else:
                            self.add_whatsapp_credits(1)
                    else:
                        self.log_message("⚠️ No WhatsApp credits for resend.")

                status_str = self._build_status_string(new_tg_ok, new_wa_ok, new_wa_skip)

                self.root.after(0, lambda: self.tree.item(item_id, values=(patient_id, patient_name, accession_no, os.path.basename(file_path), status_str)))

                all_ok = (not tg_enabled or new_tg_ok) and (not wa_enabled or new_wa_ok or new_wa_skip)
                if all_ok:
                    try:
                        if os.path.exists(file_path):
                            os.remove(file_path)
                            self.remove_from_index(file_path)
                            self.log_message(f"🗑️ Deleted DICOM after successful resend: {os.path.basename(file_path)}")
                    except Exception as e:
                        self.log_message(f"⚠️ Could not delete DICOM: {e}")
                    self.save_pdf_index(patient_id, patient_name, accession_no, pdf_path)
                else:
                    self.log_message(f"⚠️ Resend partial success: {status_str}")

            except Exception as e:
                self.log_message(f"❌ Resend error: {e}")
                self.root.after(0, lambda: self.tree.item(item_id, values=(row_values[0], row_values[1], row_values[2], os.path.basename(file_path), f"Error: {str(e)[:30]}")))

    # --- Clear Exams (delete Inbox and PDFs) ---
    def clear_exams_action(self):
        inbox_dir = self.config.get("receive_folder", DEFAULT_INBOX)
        pdf_dir = self.get_pdf_folder()
        res = messagebox.askyesno(
            "Confirm Clear Exams",
            "This will permanently delete:\n\n"
            f"• All files in the Inbox folder:\n   {inbox_dir}\n\n"
            f"• All saved PDF reports in:\n   {pdf_dir}\n\n"
            "Continue?"
        )
        if not res:
            return

        deleted = 0
        errors = 0
        for folder in (inbox_dir, pdf_dir):
            if folder and os.path.exists(folder):
                for fname in os.listdir(folder):
                    fpath = os.path.join(folder, fname)
                    if os.path.isfile(fpath):
                        try:
                            os.remove(fpath)
                            deleted += 1
                        except Exception as e:
                            errors += 1
                            self.log_message(f"⚠️ Could not delete {fpath}: {e}")

        try:
            self.clear_pdf_index_table()
        except Exception as e:
            self.log_message(f"⚠️ Could not clear PDF index table: {e}")
        try:
            conn = sqlite3.connect(DATABASE_PATH)
            c = conn.cursor()
            inbox_dir_norm = os.path.normpath(inbox_dir)
            c.execute("DELETE FROM dicom_index WHERE file_path LIKE ?", (f"{inbox_dir_norm}%",))
            conn.commit()
            conn.close()
        except Exception as e:
            self.log_message(f"⚠️ Could not clear dicom_index: {e}")

        to_remove = []
        for item in self.tree.get_children():
            to_remove.append(item)
        for item in to_remove:
            self.tree.delete(item)
        self.queue_data.clear()

        msg = f"Deleted {deleted} file(s) from Inbox and PDF folders."
        if errors:
            msg += f"\n{errors} file(s) could not be deleted (check console)."
        messagebox.showinfo("Clear Exams", msg)
        self.log_message(f"🧹 Clear Exams: deleted {deleted} file(s), {errors} error(s).")

    # ---------- Config Tab ----------
    def build_config_tab(self, parent):
        content_frame = Frame(parent, bg=self.bg_card)
        content_frame.pack(fill="both", expand=True)

        Label(content_frame, text="SYSTEM INITIALIZATION TARGETS", font=("Arial", 14, "bold"), fg=self.accent_cyan, bg=self.bg_card).pack(pady=15)

        form = Frame(content_frame, bg=self.bg_card)
        form.pack(padx=30, fill="both", expand=True, pady=10)

        def make_entry(lbl_txt, attr):
            f = Frame(form, bg=self.bg_card)
            f.pack(fill="x", pady=6, padx=20)
            Label(f, text=lbl_txt, font=("Arial", 9, "bold"), fg=self.text_light, bg=self.bg_card, width=25, anchor="w").pack(side="left")
            e = Entry(f, font=("Arial", 10), bg=self.bg_dark, fg=self.text_light, bd=1, insertbackground="white")
            e.pack(side="left", fill="x", expand=True, padx=5)
            e.insert(0, self.config.get(attr, ""))
            return e

        self.ent_inst_name = make_entry("Institute Name (PDF Title):", "institute_name")
        self.ent_telegram_token = make_entry("Telegram Bot Token:", "telegram_bot_token")
        self.ent_wa_key = make_entry("WhatsApp Cloud API Key (Access Token):", "whatsapp_api_key")
        self.ent_wa_phone_id = make_entry("WhatsApp Phone Number ID:", "whatsapp_phone_number_id")
        self.ent_wa_sender = make_entry("WhatsApp Sender Phone (optional):", "whatsapp_sender_phone")
        self.ent_ae_title = make_entry("Storage AE Title:", "ae_title")
        self.ent_ip_addr = make_entry("Host Local IP Address:", "ip_address")
        self.ent_port_num = make_entry("Server Dynamic Port:", "port")
        self.ent_batch_wait = make_entry("Combine Images Wait Time (sec):", "batch_wait_seconds")

        # Reset IP & Port button
        reset_frame = Frame(form, bg=self.bg_card)
        reset_frame.pack(fill="x", pady=5, padx=20)
        Button(reset_frame, text="🔄 Reset IP & Port (Auto)", font=("Arial", 9, "bold"), bg=self.accent_cyan, fg=self.bg_dark, bd=0, padx=10, pady=5, cursor="hand2", command=self.reset_ip_port).pack(anchor="w")

        f_pdf = Frame(form, bg=self.bg_card)
        f_pdf.pack(fill="x", pady=6, padx=20)
        Label(f_pdf, text="PDF Output Folder:", font=("Arial", 9, "bold"), fg=self.text_light, bg=self.bg_card, width=25, anchor="w").pack(side="left")
        self.ent_pdf_folder = Entry(f_pdf, font=("Arial", 10), bg=self.bg_dark, fg=self.text_light, bd=1, insertbackground="white")
        self.ent_pdf_folder.pack(side="left", fill="x", expand=True, padx=5)
        self.ent_pdf_folder.insert(0, self.config.get("pdf_output_folder", ""))
        Button(f_pdf, text="Browse", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0,
               command=lambda: self.pick_directory("pdf_output_folder", self.ent_pdf_folder)).pack(side="left", padx=2)

        f_dir1 = Frame(form, bg=self.bg_card)
        f_dir1.pack(fill="x", pady=6, padx=20)
        Label(f_dir1, text="Inbox Folder (temporary):", font=("Arial", 9, "bold"), fg=self.text_light, bg=self.bg_card, width=25, anchor="w").pack(side="left")
        self.ent_folder_path = Entry(f_dir1, font=("Arial", 10), bg=self.bg_dark, fg=self.text_light, bd=1, insertbackground="white")
        self.ent_folder_path.pack(side="left", fill="x", expand=True, padx=5)
        self.ent_folder_path.insert(0, self.config["receive_folder"])
        Button(f_dir1, text="Browse", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0, command=lambda: self.pick_directory("receive_folder", self.ent_folder_path)).pack(side="left", padx=2)

        f_dir2 = Frame(form, bg=self.bg_card)
        f_dir2.pack(fill="x", pady=6, padx=20)
        Label(f_dir2, text="Archive Folder (external only):", font=("Arial", 9, "bold"), fg=self.text_light, bg=self.bg_card, width=25, anchor="w").pack(side="left")
        self.ent_archive_path = Entry(f_dir2, font=("Arial", 10), bg=self.bg_dark, fg=self.text_light, bd=1, insertbackground="white")
        self.ent_archive_path.pack(side="left", fill="x", expand=True, padx=5)
        self.ent_archive_path.insert(0, self.config.get("archive_folder", DEFAULT_ARCHIVE))
        Button(f_dir2, text="Browse", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0, command=lambda: self.pick_directory("archive_folder", self.ent_archive_path)).pack(side="left", padx=2)

        self.auto_start_var = IntVar(value=1 if self.config.get("auto_start", False) else 0)
        def toggle_auto_start():
            enable = bool(self.auto_start_var.get())
            self.config["auto_start"] = enable
            if enable:
                self._add_to_startup()
            else:
                self._remove_from_startup()
            self.save_configuration()
            self.log_message(f"Auto‑start {'enabled' if enable else 'disabled'}")
        Checkbutton(form, text="🚀 Launch on System Startup (Windows)", variable=self.auto_start_var,
                    command=toggle_auto_start, bg=self.bg_card, fg=self.text_light,
                    selectcolor=self.bg_card, font=("Arial", 10, "bold")).pack(anchor="w", padx=20, pady=8)

        db_frame = Frame(form, bg=self.bg_card)
        db_frame.pack(fill="x", pady=5, padx=20)
        Label(db_frame, text="Database Path (fixed):", font=("Arial", 9, "bold"), fg=self.text_light, bg=self.bg_card).pack(anchor="w")
        lbl_db_path = Label(db_frame, text=DATABASE_PATH, font=("Arial", 9), fg="#9ca3af", bg=self.bg_card, wraplength=500, justify="left")
        lbl_db_path.pack(anchor="w", side="left")
        Button(db_frame, text="📋 Copy Path", font=("Arial", 8, "bold"), bg="#4b5563", fg=self.text_light, bd=0, padx=5, pady=2, cursor="hand2",
               command=lambda: self.copy_to_clipboard(DATABASE_PATH)).pack(side="left", padx=10)

        self.lbl_index_progress_config = Label(form, text="✅ Indexing ready.", font=("Arial", 9), fg=self.accent_green, bg=self.bg_card)
        self.lbl_index_progress_config.pack(anchor="w", padx=20, pady=(5,10))

        # Button frame for index operations
        index_btn_frame = Frame(form, bg=self.bg_card)
        index_btn_frame.pack(anchor="w", padx=20, pady=5)

        Button(index_btn_frame, text="🔄 Re-index Archive", font=("Arial", 9, "bold"), bg=self.accent_cyan, fg=self.bg_dark, bd=0, padx=10, pady=5, cursor="hand2", command=self.reindex_archive).pack(side="left", padx=5)

        Button(index_btn_frame, text="🗑️ Clear dicom_index", font=("Arial", 9, "bold"), bg="#ef4444", fg=self.text_light, bd=0, padx=10, pady=5, cursor="hand2", command=self.clear_dicom_index).pack(side="left", padx=5)

        Button(index_btn_frame, text="🗑️ Clear pdf_index", font=("Arial", 9, "bold"), bg="#ef4444", fg=self.text_light, bd=0, padx=10, pady=5, cursor="hand2", command=self.clear_pdf_index).pack(side="left", padx=5)

        Button(content_frame, text="Apply Node Topology Changes", font=("Arial", 11, "bold"), bg=self.accent_green, fg=self.bg_dark, width=28, bd=0, cursor="hand2", command=self.apply_and_save_node_settings).pack(pady=15)
        Label(content_frame, text="Made with ❤️ by Sandeep", font=("Arial", 9, "bold", "italic"), fg="#6b7280", bg=self.bg_card).pack(side="bottom", pady=5)

        self.lock_frame = Frame(parent, bg=self.bg_dark)
        self.lock_frame.place(relx=0, rely=0, relwidth=1, relheight=1)

        lock_card = Frame(self.lock_frame, bg=self.bg_card, bd=2, relief="ridge")
        lock_card.place(relx=0.5, rely=0.5, anchor="center", width=350, height=180)
        Label(lock_card, text="🔒 Config Control Locked", font=("Arial", 14, "bold"), fg=self.accent_cyan, bg=self.bg_card).pack(pady=(20,10))
        Label(lock_card, text="Enter password to access settings:", font=("Arial", 10), fg=self.text_light, bg=self.bg_card).pack(pady=5)
        self.config_pass_var = StringVar()
        entry_cfg_pass = Entry(lock_card, textvariable=self.config_pass_var, show="*", font=("Arial", 12), width=20, justify="center", bg=self.bg_dark, fg=self.text_light, bd=1)
        entry_cfg_pass.pack(pady=5)
        entry_cfg_pass.focus()
        btn_unlock = Button(lock_card, text="Unlock", font=("Arial", 10, "bold"), bg=self.accent_cyan, fg=self.bg_dark, width=12, command=self.unlock_config_tab)
        btn_unlock.pack(pady=10)
        entry_cfg_pass.bind("<Return>", lambda e: self.unlock_config_tab())

    # ---------- Reset IP & Port ----------
    def reset_ip_port(self):
        """Reset IP to auto-detected local IP and port to default 11112."""
        new_ip = self.get_local_ip()
        new_port = "11112"
        self.config["ip_address"] = new_ip
        self.config["port"] = new_port
        self.save_configuration()
        if self.ent_ip_addr:
            self.ent_ip_addr.delete(0, END)
            self.ent_ip_addr.insert(0, new_ip)
        if self.ent_port_num:
            self.ent_port_num.delete(0, END)
            self.ent_port_num.insert(0, new_port)
        if self.lbl_ip:
            self.lbl_ip.config(text=new_ip)
        if self.lbl_port:
            self.lbl_port.config(text=new_port)
        self.log_message(f"🔄 IP reset to {new_ip}, port reset to {new_port}")
        messagebox.showinfo("Reset Complete", f"IP set to {new_ip}\nPort set to {new_port}")

    # ---------- Helpers ----------
    def copy_log(self):
        if self.log_widget:
            content = self.log_widget.get("1.0", END)
            self.root.clipboard_clear()
            self.root.clipboard_append(content)
            messagebox.showinfo("Copied", "Console log copied to clipboard.")

    def clear_log(self):
        if self.log_widget:
            self.log_widget.delete("1.0", END)

    def copy_to_clipboard(self, text):
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        messagebox.showinfo("Copied", f"Path copied to clipboard:\n{text}")

    def reindex_archive(self):
        if self.indexing_in_progress:
            messagebox.showinfo("Info", "Indexing already in progress. Please wait.")
            return
        res = messagebox.askyesno("Confirm Re-index", "This will rebuild the entire dicom_index from the Archive folder. Continue?")
        if res:
            threading.Thread(target=self.index_all_existing_files, args=(True,), daemon=True).start()

    def clear_dicom_index(self):
        res = messagebox.askyesno("Confirm Clear", "This will delete ALL records from dicom_index table and VACUUM the database. Continue?")
        if not res:
            return
        try:
            conn = sqlite3.connect(DATABASE_PATH)
            c = conn.cursor()
            c.execute("DELETE FROM dicom_index")
            conn.commit()
            c.execute("VACUUM")
            conn.commit()
            conn.close()
            self.log_message("✅ dicom_index cleared and database vacuumed.")
            messagebox.showinfo("Success", "dicom_index table cleared and VACUUM completed.")
        except Exception as e:
            self.log_message(f"❌ Failed to clear dicom_index: {e}")
            messagebox.showerror("Error", f"Failed to clear dicom_index: {e}")

    def clear_pdf_index(self):
        res = messagebox.askyesno("Confirm Clear", "This will delete ALL records from pdf_index table and VACUUM the database. Continue?")
        if not res:
            return
        try:
            conn = sqlite3.connect(DATABASE_PATH)
            c = conn.cursor()
            c.execute("DELETE FROM pdf_index")
            conn.commit()
            c.execute("VACUUM")
            conn.commit()
            conn.close()
            self.log_message("✅ pdf_index cleared and database vacuumed.")
            messagebox.showinfo("Success", "pdf_index table cleared and VACUUM completed.")
        except Exception as e:
            self.log_message(f"❌ Failed to clear pdf_index: {e}")
            messagebox.showerror("Error", f"Failed to clear pdf_index: {e}")

    def sync_archive_folder_to_dashboard(self):
        archive_dir = self.config.get("archive_folder", DEFAULT_ARCHIVE)
        if not os.path.exists(archive_dir):
            return
        def worker():
            for file in os.listdir(archive_dir):
                if file.lower().endswith(".dcm"):
                    full_path = os.path.join(archive_dir, file)
                    try:
                        ds = pydicom.dcmread(full_path, stop_before_pixels=True)
                        patient_id = self._normalize_string(ds.get("PatientID", "N/A"))
                        patient_name = self._normalize_string(ds.get("PatientName", "N/A"))
                        accession_no = self._normalize_string(ds.get("AccessionNumber", "NO_ACC"))
                        self.root.after(0, lambda p=patient_id, n=patient_name, a=accession_no, f=file:
                                        self.upsert_grid_record(full_path, p, n, a, "Archived (External) 📁"))
                    except Exception:
                        continue
        threading.Thread(target=worker, daemon=True).start()

    def pick_directory(self, config_key, entry_widget):
        selected_dir = filedialog.askdirectory()
        if selected_dir:
            entry_widget.delete(0, "end")
            entry_widget.insert(0, os.path.normpath(selected_dir))

    def apply_and_save_node_settings(self):
        self.stop_folder_monitor()
        self.config["institute_name"] = self.ent_inst_name.get().strip().upper()
        new_token = self.ent_telegram_token.get().strip()
        if new_token and new_token != self.TELEGRAM_BOT_TOKEN:
            self.TELEGRAM_BOT_TOKEN = new_token
            self.config["telegram_bot_token"] = new_token
            self.last_update_id = 0
            self.update_bot_username()
            self.start_telegram_bot_polling()
        self.config["whatsapp_api_key"] = self.ent_wa_key.get().strip()
        self.config["whatsapp_phone_number_id"] = self.ent_wa_phone_id.get().strip()
        self.config["whatsapp_sender_phone"] = self.ent_wa_sender.get().strip()
        self.config["ae_title"] = self.ent_ae_title.get().strip()
        self.config["ip_address"] = self.ent_ip_addr.get().strip()
        self.config["port"] = self.ent_port_num.get().strip()
        self.config["receive_folder"] = self.ent_folder_path.get().strip()
        self.config["archive_folder"] = self.ent_archive_path.get().strip()
        self.config["pdf_output_folder"] = self.ent_pdf_folder.get().strip()
        try:
            self.config["batch_wait_seconds"] = max(1, float(self.ent_batch_wait.get().strip()))
        except (ValueError, AttributeError):
            self.config["batch_wait_seconds"] = DEFAULT_BATCH_WAIT_SECONDS
        self.save_configuration()
        messagebox.showinfo("System Config", "RAD-XR Core configurations updated successfully!")
        self.show_main_dashboard()
        self.sync_archive_folder_to_dashboard()
        threading.Thread(target=self.index_all_existing_files, daemon=True).start()
        self.start_folder_monitor()

    def manual_file_upload_trigger(self):
        file_path = filedialog.askopenfilename(filetypes=[("DICOM Files", "*.dcm"), ("All Files", "*.*")])
        if file_path:
            th = threading.Thread(target=self.autonomous_processing_pipeline, args=(file_path, True), daemon=True)
            th.start()

    def on_grid_row_double_click_resend(self, event):
        item_id = self.tree.selection()
        if not item_id:
            return
        row_values = self.tree.item(item_id, 'values')
        status = row_values[4]
        if "❌" in status:
            res = messagebox.askyesno("Resend", f"Resend this file?\n{row_values[3]}")
            if res:
                file_path = None
                for fpath, rid in self.queue_data.items():
                    if rid == item_id:
                        file_path = fpath
                        break
                if not file_path:
                    inbox_file = os.path.join(self.config["receive_folder"], row_values[3])
                    if os.path.exists(inbox_file):
                        file_path = inbox_file
                if file_path and os.path.exists(file_path):
                    self.tree.item(item_id, values=(row_values[0], row_values[1], row_values[2], row_values[3], "⚡ Resending..."))
                    th = threading.Thread(target=self.resend_single_file, args=(file_path, item_id), daemon=True)
                    th.start()
                else:
                    messagebox.showerror("Error", "File not found.")

    # ---------- DICOM Server ----------
    def is_valid_local_ip(self, ip):
        if ip == "0.0.0.0":
            return True
        try:
            local_ips = socket.gethostbyname_ex(socket.gethostname())[2]
            return ip in local_ips
        except:
            return False

    def toggle_server_process(self):
        if not self.is_listening:
            port = int(self.config["port"])
            ip = self.config["ip_address"].strip()
            if not ip:
                ip = "0.0.0.0"
            if not self.is_valid_local_ip(ip):
                new_ip = self.get_local_ip()
                self.log_message(f"⚠️ IP '{ip}' not found on any interface. Falling back to {new_ip}")
                ip = new_ip
                self.config["ip_address"] = ip
                self.save_configuration()
                if self.ent_ip_addr:
                    self.ent_ip_addr.delete(0, END)
                    self.ent_ip_addr.insert(0, ip)
                if self.lbl_ip:
                    self.lbl_ip.config(text=ip)
            try:
                test_sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                test_sock.settimeout(2)
                result = test_sock.connect_ex((ip, port))
                test_sock.close()
                if result == 0:
                    self.log_message(f"❌ Port {port} is already in use on {ip}.")
                    messagebox.showerror("Port Error", f"Port {port} already in use.")
                    return
            except Exception as e:
                self.log_message(f"⚠️ Could not check port: {e}")

            try:
                os.makedirs(self.config["receive_folder"], exist_ok=True)
                os.makedirs(self.config["archive_folder"], exist_ok=True)
            except Exception as e:
                self.log_message(f"❌ Failed to create folders: {e}")
                messagebox.showerror("Folder Error", f"Cannot create folders:\n{e}")
                return

            self.is_listening = True
            self.status_var.set("● Starting...")
            self.lbl_status_indicator.config(fg=self.accent_cyan)
            self.btn_toggle_server.config(text="Starting...", bg="#4b5563")
            self.log_message("🚀 Starting DICOM server...")
            self.log_message(f"   AE Title: {self.config['ae_title']}")
            self.log_message(f"   IP: {ip}")
            self.log_message(f"   Port: {port}")
            self.log_message(f"   Inbox: {self.config['receive_folder']}")
            self.server_thread = threading.Thread(target=self.run_dicom_scp_listener, args=(ip,), daemon=True)
            self.server_thread.start()
        else:
            self.is_listening = False
            if self.server_instance:
                self.server_instance.shutdown()
            self.status_var.set("● Stopped")
            self.lbl_status_indicator.config(fg=self.accent_red)
            self.btn_toggle_server.config(text="Start Server", bg=self.accent_green)
            self.log_message("⏹️ Server stopped by user.")

    def run_dicom_scp_listener(self, bind_ip):
        ae = AE()
        raw_ae = self.config["ae_title"]
        sanitized = ''.join(c if c.isalnum() or c in ' _' else '_' for c in raw_ae)
        ae.ae_title = sanitized
        self.log_message(f"   Sanitized AE Title: '{ae.ae_title}'")
        try:
            ae.add_supported_context("1.2.840.10008.1.1")
            for context in AllStoragePresentationContexts:
                ae.add_supported_context(
                    context.abstract_syntax,
                    ACCEPTED_TRANSFER_SYNTAXES
                )
            self.log_message(f"Loaded {len(AllStoragePresentationContexts)} Storage Presentation Contexts.")
            handlers = [
                (evt.EVT_C_STORE, self.handle_incoming_c_store),
                (evt.EVT_C_ECHO, self.handle_incoming_c_echo)
            ]
            self.log_message(f"⏳ Attempting to bind to {bind_ip}:{self.config['port']}")
            self.server_instance = ae.start_server(
                (bind_ip, int(self.config["port"])),
                block=False,
                evt_handlers=handlers
            )
            self.log_message("✅ start_server() returned successfully.")
            time.sleep(2)
            test_sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            test_sock.settimeout(3)
            try:
                test_sock.connect((bind_ip, int(self.config["port"])))
                test_sock.close()
                self.root.after(0, self._server_started_successfully)
                self.log_message(f"✅ DICOM server is LISTENING on {bind_ip}:{self.config['port']}")
                while self.is_listening:
                    time.sleep(0.5)
            except Exception as conn_err:
                self.log_message(f"❌ Verification connection failed: {conn_err}")
                self.root.after(0, self._server_failed_to_start, f"Port {self.config['port']} not open.")
                return
        except Exception as e:
            self.log_message(f"❌ Server start exception: {e}")
            self.root.after(0, self._server_failed_to_start, str(e))

    def _server_started_successfully(self):
        self.status_var.set("● Listening")
        self.lbl_status_indicator.config(fg=self.accent_green)
        self.btn_toggle_server.config(text="Stop Server", bg=self.accent_red)
        self.log_message("🎉 Server is now active and accepting connections.")

    def _server_failed_to_start(self, error_msg):
        self.is_listening = False
        self.status_var.set("● Stopped")
        self.lbl_status_indicator.config(fg=self.accent_red)
        self.btn_toggle_server.config(text="Start Server", bg=self.accent_green)
        self.log_message(f"❌ Server failed: {error_msg}")
        messagebox.showerror("Server Error", f"Failed to start DICOM server:\n{error_msg}")

    def handle_incoming_c_echo(self, event):
        self.log_message(f"✅ C-ECHO received from {event.assoc.requestor.ae_title}")
        return 0x0000

    def handle_incoming_c_store(self, event):
        try:
            ds = event.dataset
            ds.file_meta = event.file_meta
            accession_number = self._normalize_string(ds.get("AccessionNumber", "UNKNOWN_ACC"))
            sop_instance_uid = self._normalize_string(ds.get("SOPInstanceUID", ""))
            unique_suffix = sop_instance_uid[-12:] if sop_instance_uid else f"{int(time.time() * 1000)}"
            filename = f"RADXR_{accession_number}_{unique_suffix}.dcm"
            filepath = os.path.join(self.config["receive_folder"], filename)
            ds.save_as(filepath, write_like_original=False)
            self.log_message(f"📥 C-STORE received from {event.assoc.requestor.ae_title}")
            threading.Thread(
                target=self.queue_file_for_patient_batch,
                args=(filepath, False),
                daemon=True
            ).start()
            return 0x0000
        except Exception as e:
            self.log_message(f"❌ C-STORE error: {e}")
            return 0xC000

    # ---------- PDF Generation ----------
    def generate_pdf_report_from_dicom(self, dcm_paths, output_pdf_path):
        """
        Builds ONE PDF report from one or more DICOM files.
        Each image (or frame) is placed on its own page.
        Returns (patient_id, patient_name, accession_no, study_date).
        """
        from PIL import Image as PILImage

        if isinstance(dcm_paths, (str, bytes, os.PathLike)):
            dcm_paths = [dcm_paths]

        first_ds = pydicom.dcmread(dcm_paths[0], stop_before_pixels=True)
        patient_id = self._normalize_string(first_ds.get("PatientID", "N/A"))
        patient_name = self._normalize_string(first_ds.get("PatientName", "N/A"))
        accession_no = self._normalize_string(first_ds.get("AccessionNumber", "NO_ACC"))
        study_date = self._normalize_string(first_ds.get("StudyDate", ""))

        c = canvas.Canvas(output_pdf_path, pagesize=letter)
        width, height = letter

        top_margin = 8 * 72 / 25.4
        margin_lr = 3 * 72 / 25.4

        metadata = [
            ("Patient Name", patient_name),
            ("Patient ID", patient_id),
            ("Patient Sex", self._normalize_string(first_ds.get("PatientSex", "N/A"))),
            ("Study Date", study_date if study_date else "N/A"),
            ("Modality", self._normalize_string(first_ds.get("Modality", "N/A"))),
            ("Accession No", accession_no)
        ]
        available_metadata = [(k, v) for k, v in metadata if v.strip() and v != "N/A"]

        footer_image_path = self.pdf_footer_image if self.pdf_footer_image and os.path.exists(self.pdf_footer_image) else None
        footer_img_w = 0
        footer_img_h = 0
        if footer_image_path:
            try:
                footer_img = PILImage.open(footer_image_path)
                f_img_w, f_img_h = footer_img.size
                draw_w = width
                draw_h = (f_img_h / f_img_w) * draw_w
                max_footer_h = 144
                if draw_h > max_footer_h:
                    draw_h = max_footer_h
                    draw_w = (f_img_w / f_img_h) * draw_h
                footer_img_w = draw_w
                footer_img_h = draw_h
                footer_img.close()
            except Exception as e:
                self.log_message(f"Footer image error: {e}")
                footer_img_w = 0
                footer_img_h = 0

        total_pages = 0
        for dcm_path in dcm_paths:
            try:
                ds_probe = pydicom.dcmread(dcm_path, stop_before_pixels=True)
                n_frames = int(getattr(ds_probe, "NumberOfFrames", 1) or 1)
            except Exception:
                n_frames = 1
            total_pages += n_frames
        if total_pages == 0:
            total_pages = len(dcm_paths)

        page_counter = 0
        temp_files = []

        for dcm_path in dcm_paths:
            try:
                ds = pydicom.dcmread(dcm_path)
                pixel_array = self._safe_load_pixel_array(ds)
            except Exception as e:
                self.log_message(f"❌ Skipping unreadable DICOM {os.path.basename(str(dcm_path))}: {e}")
                continue

            num_frames = 1
            if hasattr(ds, "NumberOfFrames") and ds.NumberOfFrames > 1:
                num_frames = int(ds.NumberOfFrames)
            elif len(pixel_array.shape) == 3 and pixel_array.shape[0] < pixel_array.shape[1]:
                num_frames = pixel_array.shape[0]

            for frame_idx in range(num_frames):
                page_counter += 1
                frame_array = pixel_array[frame_idx] if num_frames > 1 else pixel_array

                if frame_array.dtype != np.uint8:
                    p_min = frame_array.min()
                    p_max = frame_array.max()
                    if p_max > p_min:
                        frame_array = (((frame_array - p_min) / (p_max - p_min)) * 255).astype(np.uint8)
                    else:
                        frame_array = frame_array.astype(np.uint8)

                image = PILImage.fromarray(frame_array)
                if image.mode != "RGB":
                    image = image.convert("RGB")

                img_w, img_h = image.size

                temp_img_path = f"workflow_temp_frame_{page_counter}_{int(time.time())}.jpg"
                image.save(temp_img_path, quality=95)
                temp_files.append(temp_img_path)
                image.close()
                del image
                gc.collect()

                header_y = height - top_margin
                c.setFont("Helvetica-Bold", 14)
                c.drawString(margin_lr, header_y, self.config["institute_name"])
                c.setFont("Helvetica-Oblique", 9)
                c.drawRightString(width - margin_lr, header_y, f"Page {page_counter} of {total_pages}")

                c.setLineWidth(1)
                c.setStrokeColorRGB(0.1, 0.5, 0.7)
                c.line(margin_lr, header_y - 6, width - margin_lr, header_y - 6)

                y_meta_start = header_y - 20
                y_text = y_meta_start
                col = 0
                for label, value in available_metadata:
                    x_pos = margin_lr if col == 0 else width / 2 + margin_lr - 40
                    c.setFont("Helvetica-Bold", 9)
                    c.drawString(x_pos, y_text, f"{label}: ")
                    c.setFont("Helvetica", 9)
                    c.drawString(x_pos + 80, y_text, value)
                    if col == 1:
                        y_text -= 15
                        col = 0
                    else:
                        col = 1
                if col == 1:
                    y_text -= 15

                c.setLineWidth(0.5)
                c.line(margin_lr, y_text, width - margin_lr, y_text)
                main_image_top = y_text - 15

                if footer_img_h > 0:
                    gap = 5
                    main_image_bottom = footer_img_h + gap
                else:
                    main_image_bottom = 0

                avail_height = main_image_top - main_image_bottom
                avail_width = width - 2 * margin_lr
                scale = 1.0
                if img_w > avail_width or img_h > avail_height:
                    scale = min(avail_width / img_w, avail_height / img_h)
                draw_main_w = img_w * scale
                draw_main_h = img_h * scale
                x_main = (width - draw_main_w) / 2
                y_main = main_image_bottom + (avail_height - draw_main_h) / 2
                c.drawImage(temp_img_path, x_main, y_main, width=draw_main_w, height=draw_main_h,
                            preserveAspectRatio=True, anchor='c')

                if footer_img_w > 0 and footer_img_h > 0:
                    c.drawImage(footer_image_path, 0, 0, width=footer_img_w, height=footer_img_h,
                                preserveAspectRatio=False, anchor='sw')

                if page_counter < total_pages:
                    c.showPage()

        c.save()

        for temp_path in temp_files:
            if os.path.exists(temp_path):
                self._delete_file_with_retry(temp_path, max_retries=5, delay=0.2)

        gc.collect()

        return patient_id, patient_name, accession_no, study_date

    def _delete_file_with_retry(self, file_path, max_retries=5, delay=0.2):
        for attempt in range(max_retries):
            try:
                if os.path.exists(file_path):
                    os.remove(file_path)
                return
            except Exception as e:
                if attempt < max_retries - 1:
                    time.sleep(delay * (attempt + 1))
                else:
                    self.log_message(f"⚠️ Could not delete {file_path} after {max_retries} attempts: {e}")

    # ---------- Internet Check ----------
    def _check_internet(self):
        try:
            sock.create_connection(("8.8.8.8", 53), timeout=3)
            return True
        except OSError:
            return False

    def _log_no_internet(self):
        if self.log_widget:
            self.log_widget.insert(END, "\n\n")
            self.log_widget.insert(END, "!!! NO INTERNET CONNECTION !!!\n", ("big_red",))
            self.log_widget.insert(END, "Telegram send will fail.\n", ("big_red",))
            self.log_widget.tag_configure("big_red", foreground="red", font=("Arial", 16, "bold"))
            self.log_widget.see(END)

    # ---------- Processing Pipeline ----------
    def autonomous_processing_pipeline(self, dcm_path, is_manual_import=False):
        with self.processing_lock:
            pdf_output_path = ""
            try:
                ds = pydicom.dcmread(dcm_path)
                patient_id = self._normalize_string(ds.get("PatientID", "N/A"))
                patient_name = self._normalize_string(ds.get("PatientName", "N/A"))
                accession_no = self._normalize_string(ds.get("AccessionNumber", "UNKNOWN"))
                study_date = self._normalize_string(ds.get("StudyDate", ""))

                self.index_dicom_file(dcm_path, patient_id, patient_name, accession_no)

                pdf_output_path = os.path.join(
                    self.get_pdf_folder(),
                    self.build_pdf_filename(patient_name, study_date)
                )
                self.generate_pdf_report_from_dicom([dcm_path], pdf_output_path)

                file_key = dcm_path
                self.root.after(0, lambda: self.upsert_grid_record(file_key, patient_id, patient_name, accession_no, "⏳ Processing"))
                self.root.after(0, lambda: self.upsert_grid_record(file_key, patient_id, patient_name, accession_no, "📤 Sending"))

                tg_enabled = bool(self.TELEGRAM_BOT_TOKEN)
                wa_enabled = bool(self.config.get("whatsapp_api_key") and self.config.get("whatsapp_phone_number_id"))

                tg_ok = False
                wa_ok = False
                wa_skip = False

                accession_missing = not accession_no or accession_no in ("N/A", "UNKNOWN")

                if tg_enabled:
                    if not self._check_internet():
                        self._log_no_internet()
                        tg_ok = False
                    else:
                        if self.decrement_telegram_credits():
                            tg_ok = self.send_to_all_telegram(pdf_output_path, patient_id, patient_name, accession_no)
                            if not tg_ok:
                                self.add_telegram_credits(1)
                                self.log_message("⚠️ Telegram send failed, credit refunded.")
                        else:
                            self.log_message("⚠️ Insufficient Telegram credits.")

                if wa_enabled:
                    if accession_missing:
                        wa_skip = True
                        self.log_message(f"ℹ️ WhatsApp skipped: Accession Number missing for {patient_name} ({patient_id})")
                    else:
                        if self.decrement_whatsapp_credits():
                            wa_ok = self.dispatch_to_whatsapp_business(pdf_output_path, patient_id, patient_name, accession_no)
                            if not wa_ok:
                                self.add_whatsapp_credits(1)
                                self.log_message("⚠️ WhatsApp send failed, credit refunded.")
                        else:
                            self.log_message("⚠️ Insufficient WhatsApp credits.")

                status_parts = []
                if tg_enabled:
                    status_parts.append(f"Telegram {'✅' if tg_ok else '❌'}")
                if wa_enabled:
                    if wa_skip:
                        status_parts.append("WhatsApp ⏭️")
                    else:
                        status_parts.append(f"WhatsApp {'✅' if wa_ok else '❌'}")
                if not status_parts:
                    status_parts.append("No platforms configured")
                status_str = ", ".join(status_parts)

                self.root.after(0, lambda: self.upsert_grid_record(file_key, patient_id, patient_name, accession_no, status_str))

                all_ok = (not tg_enabled or tg_ok) and (not wa_enabled or wa_ok or wa_skip)
                if all_ok:
                    self.save_pdf_index(patient_id, patient_name, accession_no, pdf_output_path)
                    self.log_message(f"💾 PDF saved and indexed: {os.path.basename(pdf_output_path)}")
                    try:
                        if os.path.exists(dcm_path):
                            os.remove(dcm_path)
                            self.remove_from_index(dcm_path)
                            self.log_message(f"🗑️ Deleted DICOM after successful send: {os.path.basename(dcm_path)}")
                    except Exception as e:
                        self.log_message(f"⚠️ Could not delete DICOM: {e}")
                else:
                    self.log_message(f"⚠️ Not all platforms succeeded, DICOM kept: {os.path.basename(dcm_path)}")

            except Exception as e:
                self.root.after(0, lambda: messagebox.showerror("Pipeline Error", f"Error: {str(e)}"))
                self.log_message(f"❌ Pipeline error: {e}")
                if 'file_key' in locals():
                    self.root.after(0, lambda: self.upsert_grid_record(file_key, "N/A", "N/A", "UNKNOWN", f"Error: {str(e)[:30]}"))

    # ---------- Patient Batching ----------
    def queue_file_for_patient_batch(self, dcm_path, is_manual_import=False):
        try:
            ds = pydicom.dcmread(dcm_path, stop_before_pixels=True)
            patient_id = self._normalize_string(ds.get("PatientID", "N/A"))
            patient_name = self._normalize_string(ds.get("PatientName", "N/A"))
            accession_no = self._normalize_string(ds.get("AccessionNumber", "UNKNOWN"))
        except Exception as e:
            self.log_message(f"❌ Could not read DICOM header, skipping: {e}")
            return

        self.index_dicom_file(dcm_path, patient_id, patient_name, accession_no)

        self.root.after(0, lambda: self.upsert_grid_record(
            dcm_path, patient_id, patient_name, accession_no, "📥 Received - Queued"))

        if patient_id and patient_id != "N/A":
            batch_key = f"PID::{patient_id}"
        elif patient_name and patient_name != "N/A":
            batch_key = f"NAME::{patient_name.lower()}"
        else:
            batch_key = f"UNKNOWN::{time.time_ns()}"

        try:
            wait_seconds = float(self.config.get("batch_wait_seconds", DEFAULT_BATCH_WAIT_SECONDS) or DEFAULT_BATCH_WAIT_SECONDS)
        except Exception:
            wait_seconds = DEFAULT_BATCH_WAIT_SECONDS

        with self.batch_lock:
            entry = self.pending_batches.get(batch_key)
            if entry:
                entry["items"].append((dcm_path, is_manual_import))
                if accession_no and accession_no != "UNKNOWN":
                    entry["accession_no"] = accession_no
                if entry.get("timer"):
                    entry["timer"].cancel()
            else:
                entry = {
                    "patient_id": patient_id,
                    "patient_name": patient_name,
                    "accession_no": accession_no,
                    "items": [(dcm_path, is_manual_import)],
                }
                self.pending_batches[batch_key] = entry

            timer = threading.Timer(wait_seconds, self._flush_patient_batch, args=(batch_key,))
            timer.daemon = True
            entry["timer"] = timer
            timer.start()

        self.log_message(
            f"🗂️ Queued image for {patient_name} (ID: {patient_id}) — "
            f"{len(entry['items'])} image(s) so far, waiting {wait_seconds:.0f}s..."
        )

    def _flush_patient_batch(self, batch_key):
        with self.batch_lock:
            entry = self.pending_batches.pop(batch_key, None)
        if not entry:
            return
        threading.Thread(
            target=self.process_patient_batch,
            args=(entry["items"], entry["patient_id"], entry["patient_name"], entry["accession_no"]),
            daemon=True
        ).start()

    def process_patient_batch(self, items, patient_id, patient_name, accession_no):
        with self.processing_lock:
            pdf_output_path = ""
            dcm_paths = [dcm for dcm, _ in items]
            try:
                self.log_message(
                    f"🧩 Combining {len(dcm_paths)} image(s) for {patient_name} "
                    f"(ID: {patient_id}) into one PDF..."
                )

                study_date = ""
                try:
                    first_ds = pydicom.dcmread(dcm_paths[0], stop_before_pixels=True)
                    study_date = self._normalize_string(first_ds.get("StudyDate", ""))
                except Exception:
                    pass

                for dcm in dcm_paths:
                    self.root.after(0, lambda f=dcm: self.upsert_grid_record(
                        f, patient_id, patient_name, accession_no, "⏳ Processing (Batch)"))

                pdf_output_path = os.path.join(
                    self.get_pdf_folder(),
                    self.build_pdf_filename(patient_name, study_date)
                )
                self.generate_pdf_report_from_dicom(dcm_paths, pdf_output_path)

                for dcm in dcm_paths:
                    self.root.after(0, lambda f=dcm: self.upsert_grid_record(
                        f, patient_id, patient_name, accession_no, "📤 Sending"))

                tg_enabled = bool(self.TELEGRAM_BOT_TOKEN)
                wa_enabled = bool(self.config.get("whatsapp_api_key") and self.config.get("whatsapp_phone_number_id"))

                tg_ok = False
                wa_ok = False
                wa_skip = False

                accession_missing = not accession_no or accession_no in ("N/A", "UNKNOWN")

                if tg_enabled:
                    if not self._check_internet():
                        self._log_no_internet()
                        tg_ok = False
                    else:
                        if self.decrement_telegram_credits():
                            tg_ok = self.send_to_all_telegram(pdf_output_path, patient_id, patient_name, accession_no)
                            if not tg_ok:
                                self.add_telegram_credits(1)
                                self.log_message("⚠️ Telegram send failed, credit refunded.")
                        else:
                            self.log_message("⚠️ Insufficient Telegram credits.")

                if wa_enabled:
                    if accession_missing:
                        wa_skip = True
                        self.log_message(f"ℹ️ WhatsApp skipped: Accession Number missing for {patient_name} ({patient_id})")
                    else:
                        if self.decrement_whatsapp_credits():
                            wa_ok = self.dispatch_to_whatsapp_business(pdf_output_path, patient_id, patient_name, accession_no)
                            if not wa_ok:
                                self.add_whatsapp_credits(1)
                                self.log_message("⚠️ WhatsApp send failed, credit refunded.")
                        else:
                            self.log_message("⚠️ Insufficient WhatsApp credits.")

                status_parts = []
                if tg_enabled:
                    status_parts.append(f"Telegram {'✅' if tg_ok else '❌'}")
                if wa_enabled:
                    if wa_skip:
                        status_parts.append("WhatsApp ⏭️")
                    else:
                        status_parts.append(f"WhatsApp {'✅' if wa_ok else '❌'}")
                if not status_parts:
                    status_parts.append("No platforms configured")
                status_str = ", ".join(status_parts)

                for dcm in dcm_paths:
                    self.root.after(0, lambda f=dcm: self.upsert_grid_record(
                        f, patient_id, patient_name, accession_no, status_str))

                all_ok = (not tg_enabled or tg_ok) and (not wa_enabled or wa_ok or wa_skip)
                if all_ok:
                    self.save_pdf_index(patient_id, patient_name, accession_no, pdf_output_path)
                    self.log_message(f"💾 PDF saved and indexed: {os.path.basename(pdf_output_path)}")
                    for dcm in dcm_paths:
                        try:
                            if os.path.exists(dcm):
                                os.remove(dcm)
                                self.remove_from_index(dcm)
                                self.log_message(f"🗑️ Deleted DICOM after batch send: {os.path.basename(dcm)}")
                        except Exception as e:
                            self.log_message(f"⚠️ Could not delete DICOM {dcm}: {e}")
                else:
                    self.log_message(f"⚠️ Batch not fully sent, DICOMs kept.")

            except Exception as e:
                self.root.after(0, lambda: messagebox.showerror("Pipeline Error", f"Error: {str(e)}"))
                self.log_message(f"❌ Batch pipeline error: {e}")
                for dcm in dcm_paths:
                    self.root.after(0, lambda f=dcm: self.upsert_grid_record(
                        f, patient_id, patient_name, accession_no, f"Error: {str(e)[:30]}"))

    # ---------- Grid update ----------
    def upsert_grid_record(self, file_path, p_id, p_name, acc_no, status):
        if file_path in self.queue_data:
            row_id = self.queue_data[file_path]
            self.tree.item(row_id, values=(p_id, p_name, acc_no, os.path.basename(file_path), status))
        else:
            row_id = self.tree.insert("", "end", values=(p_id, p_name, acc_no, os.path.basename(file_path), status))
            self.queue_data[file_path] = row_id

    # ---------- Caption builder ----------
    def build_beautiful_caption_string(self, p_id, p_name, acc_no, include_footer=True):
        caption = (
            f"🏥 *{self.config['institute_name']}*\n"
            f"━━━━━━━━━━━━━━━━━━━━━\n"
            f"👤 *Patient Name:* {p_name}\n"
            f"🆔 *Patient ID:* {p_id}\n"
            f"🔢 *Accession No:* {acc_no}\n"
        )
        if include_footer and self.footer_message.strip():
            caption += f"\n{self.footer_message.strip()}\n"
        caption += f"\n*Made with ❤️ by Sandeep*"
        return caption

    # ---------- Telegram send (to admin users and groups) ----------
    def send_to_all_telegram(self, file_path, p_id, p_name, acc_no):
        admin_users = self.get_admin_users()
        groups = self.get_all_groups()
        recipients = admin_users + groups

        if not recipients:
            self.log_message("No admin users or groups to send to.")
            return True

        success = True
        for chat_id in recipients:
            ok = self.dispatch_to_telegram(file_path, p_id, p_name, acc_no, chat_id)
            if not ok:
                success = False
                self.log_message(f"❌ Failed to send to {chat_id} (admin/group)")
        return success

    def dispatch_to_telegram(self, file_path, p_id, p_name, acc_no, target_chat_id):
        url = f"https://api.telegram.org/bot{self.TELEGRAM_BOT_TOKEN}/sendDocument"
        caption_text = self.build_beautiful_caption_string(p_id, p_name, acc_no, include_footer=True)

        max_attempts = 5
        base_delay = 1
        timeout = 180

        for attempt in range(1, max_attempts + 1):
            try:
                with open(file_path, "rb") as document:
                    payload = {
                        "chat_id": target_chat_id,
                        "caption": caption_text,
                        "parse_mode": "Markdown"
                    }
                    files = {"document": (os.path.basename(file_path), document, "application/pdf")}
                    res = requests.post(url, data=payload, files=files, timeout=timeout)
                    if res.status_code == 200:
                        self.log_message(f"✅ Telegram sent to {target_chat_id} (attempt {attempt})")
                        return True
                    else:
                        self.log_message(f"⚠️ Telegram attempt {attempt} failed with status {res.status_code}: {res.text}")
            except Exception as e:
                self.log_message(f"⚠️ Telegram attempt {attempt} error: {e}")

            if attempt < max_attempts:
                wait = base_delay * (2 ** (attempt - 1))
                self.log_message(f"⏳ Retrying Telegram in {wait} seconds...")
                time.sleep(wait)

        self.log_message(f"❌ Telegram send failed after {max_attempts} attempts.")
        return False

    # ---------- WhatsApp ----------
    def dispatch_to_whatsapp_business(self, file_path, p_id, p_name, acc_no):
        if not self.decrement_whatsapp_credits():
            self.log_message("❌ No WhatsApp credits available.")
            return False

        target_phone = "".join(filter(str.isdigit, acc_no))
        if len(target_phone) < 10:
            return False
        if len(target_phone) == 10:
            target_phone = "91" + target_phone

        phone_number_id = self.config["whatsapp_phone_number_id"]
        api_key = self.config["whatsapp_api_key"]
        if not phone_number_id or not api_key:
            return False

        headers = {"Authorization": f"Bearer {api_key}"}
        upload_url = f"https://graph.facebook.com/v18.0/{phone_number_id}/media"
        try:
            caption_text = self.build_beautiful_caption_string(p_id, p_name, acc_no, include_footer=True)
            clean_caption = caption_text.replace("*", "").replace("_", "")
            with open(file_path, "rb") as f:
                files = {
                    "file": (os.path.basename(file_path), f, "application/pdf"),
                    "messaging_product": (None, "whatsapp")
                }
                upload_res = requests.post(upload_url, headers=headers, files=files, timeout=60)
                if upload_res.status_code != 200:
                    self.log_message(f"WhatsApp upload failed: {upload_res.text}")
                    return False
                media_id = upload_res.json().get("id")
                if not media_id:
                    return False

                msg_url = f"https://graph.facebook.com/v18.0/{phone_number_id}/messages"
                payload = {
                    "messaging_product": "whatsapp",
                    "to": target_phone,
                    "type": "document",
                    "document": {
                        "id": media_id,
                        "filename": os.path.basename(file_path),
                        "caption": clean_caption
                    }
                }
                msg_res = requests.post(msg_url, headers=headers, json=payload, timeout=60)
                if msg_res.status_code == 200:
                    return True
                else:
                    self.log_message(f"WhatsApp send failed: {msg_res.text}")
                    return False
        except Exception as e:
            self.log_message(f"WhatsApp error: {e}")
            return False

    # ---------- Helper to send document ----------
    def _send_document(self, base_url, chat_id, file_path, caption=""):
        try:
            with open(file_path, "rb") as f:
                files = {"document": (os.path.basename(file_path), f, "application/octet-stream")}
                payload = {"chat_id": chat_id, "caption": caption}
                res = requests.post(f"{base_url}/sendDocument", data=payload, files=files, timeout=30)
                return res.status_code == 200
        except Exception as e:
            self.log_message(f"Error sending document: {e}")
            return False

    # ---------- Telegram Bot Polling (with enhanced commands) ----------
    def start_telegram_bot_polling(self):
        t = threading.Thread(target=self.telegram_bot_polling_worker, daemon=True)
        t.start()

    def telegram_bot_polling_worker(self):
        base_url = f"https://api.telegram.org/bot{self.TELEGRAM_BOT_TOKEN}"
        expecting_config = False
        config_update_message_id = None

        while self.bot_running:
            try:
                url = f"{base_url}/getUpdates?offset={self.last_update_id + 1}&timeout=10"
                response = requests.get(url, timeout=15)
                if response.status_code == 200:
                    data = response.json()
                    for update in data.get("result", []):
                        self.last_update_id = update["update_id"]
                        if "message" not in update:
                            continue
                        msg = update["message"]
                        chat_id = str(msg["chat"]["id"])
                        text = msg.get("text", "").strip()
                        user = msg.get("from", {})
                        user_id = str(user.get("id", ""))
                        username = user.get("username", "")
                        full_name = f"{user.get('first_name', '')} {user.get('last_name', '')}".strip()
                        is_group = msg["chat"]["type"] in ("group", "supergroup")

                        if user_id:
                            self.save_telegram_user(user_id, username, full_name)

                        # ---- /start ----
                        if text.lower().startswith("/start"):
                            welcome = (
                                f"🏥 *Welcome to RAD-XR Portal Search Node*\n\n"
                                f"Your User ID: `{chat_id}`\n\n"
                                "To retrieve your patient's report(s), send the following format:\n"
                                "`[PATIENT ID]`\n"
                                "`[PATIENT FIRST NAME]`\n\n"
                                "*Example:*\n"
                                "`1898`\n"
                                "`sandeep`"
                            )
                            if self.footer_message.strip():
                                welcome += f"\n\n📝 *Message from Admin:* {self.footer_message.strip()}"
                            self._send_message(base_url, chat_id, welcome)
                            continue

                        # ---- Master commands (check user_id, not chat_id) ----
                        if user_id == self.TELEGRAM_MASTER_USER_ID:
                            # ---- /newbot ----
                            if text.lower().startswith("/newbot "):
                                new_token = text[8:].strip()
                                if new_token:
                                    self.TELEGRAM_BOT_TOKEN = new_token
                                    self.last_update_id = 0
                                    self.config["telegram_bot_token"] = new_token
                                    self.save_configuration()
                                    base_url = f"https://api.telegram.org/bot{self.TELEGRAM_BOT_TOKEN}"
                                    self.update_bot_username()
                                    self.refresh_bot_info_gui()
                                    self._send_message(base_url, chat_id, "✅ Bot token updated successfully.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Provide a token: `/newbot <token>`")
                                continue

                            # ---- /adduser ----
                            if text.lower().startswith("/adduser "):
                                uid = text[9:].strip()
                                if uid:
                                    self.add_admin_user(uid, user_id)
                                    self.save_telegram_user(uid, "", "")
                                    self._send_message(base_url, chat_id, f"✅ User `{uid}` added to admin list.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Provide user ID.")
                                continue

                            # ---- /remove ----
                            if text.lower().startswith("/remove "):
                                uid = text[8:].strip()
                                if uid == self.TELEGRAM_MASTER_USER_ID:
                                    self._send_message(base_url, chat_id, "❌ Cannot remove master.")
                                elif uid:
                                    self.remove_admin_user(uid)
                                    self._send_message(base_url, chat_id, f"✅ User `{uid}` removed from admin list.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Provide user ID.")
                                continue

                            # ---- /users ----
                            if text.lower() == "/users":
                                self._send_message(base_url, chat_id, "📊 Generating user list...")
                                self._export_users_excel(base_url, chat_id)
                                continue

                            # ---- /addgroup ----
                            if text.lower().startswith("/addgroup "):
                                gid = text[10:].strip()
                                if gid:
                                    self.add_group(gid, user_id)
                                    self._send_message(base_url, chat_id, f"✅ Group `{gid}` added.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Provide group ID.")
                                continue

                            if text.lower() == "/addgroup":
                                # No argument: if in a group, auto-add this group
                                if is_group:
                                    gid = chat_id
                                    self.add_group(gid, user_id)
                                    self._send_message(base_url, chat_id, f"✅ This group `{gid}` added.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ This is a private chat. Please provide group ID: `/addgroup <group_id>`")
                                continue

                            # ---- /removegroup ----
                            if text.lower().startswith("/removegroup "):
                                gid = text[13:].strip()
                                if gid:
                                    self.remove_group(gid)
                                    self._send_message(base_url, chat_id, f"✅ Group `{gid}` removed.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Provide group ID. Use `/removegroup` without argument to list all groups.")
                                continue

                            if text.lower() == "/removegroup":
                                # No argument: list all groups
                                group_list = self._format_group_list()
                                self._send_message(base_url, chat_id, group_list)
                                continue

                            # ---- /getconfig ----
                            if text.lower() == "/getconfig":
                                if os.path.exists(CONFIG_FILE):
                                    self._send_message(base_url, chat_id, "📄 Sending current config file...")
                                    ok = self._send_document(base_url, chat_id, CONFIG_FILE, "Current config file")
                                    if not ok:
                                        self._send_message(base_url, chat_id, "❌ Failed to send config file.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Config file not found.")
                                continue

                            # ---- /getdb ----
                            if text.lower() == "/getdb":
                                if os.path.exists(DATABASE_PATH):
                                    self._send_message(base_url, chat_id, "📄 Sending database file...")
                                    ok = self._send_document(base_url, chat_id, DATABASE_PATH, "Database file (radxr_index.db)")
                                    if not ok:
                                        self._send_message(base_url, chat_id, "❌ Failed to send database file.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Database file not found.")
                                continue

                            # ---- /updateconfig ----
                            if text.lower() == "/updateconfig":
                                self._send_message(base_url, chat_id, "Please send the new config file as a document (JSON).")
                                expecting_config = True
                                config_update_message_id = msg["message_id"]
                                continue

                            # ---- /message ----
                            if text.lower().startswith("/message "):
                                new_msg = text[9:].strip()
                                self.footer_message = new_msg
                                self.config["footer_message"] = new_msg
                                self.save_configuration()
                                self._send_message(base_url, chat_id, "✅ Caption updated.")
                                continue

                            # ---- /footer ----
                            if text.lower().startswith("/footer"):
                                reply_to = msg.get("reply_to_message")
                                if reply_to and "photo" in reply_to:
                                    try:
                                        photos = reply_to["photo"]
                                        largest = photos[-1]
                                        file_id = largest["file_id"]
                                        get_file_url = f"{base_url}/getFile?file_id={file_id}"
                                        file_resp = requests.get(get_file_url, timeout=10)
                                        if file_resp.status_code == 200:
                                            file_data = file_resp.json()
                                            if file_data.get("ok"):
                                                file_path = file_data["result"]["file_path"]
                                                download_url = f"https://api.telegram.org/file/bot{self.TELEGRAM_BOT_TOKEN}/{file_path}"
                                                img_resp = requests.get(download_url, timeout=30)
                                                if img_resp.status_code == 200:
                                                    with open(FOOTER_IMAGE_PATH, "wb") as f:
                                                        f.write(img_resp.content)
                                                    self.pdf_footer_image = FOOTER_IMAGE_PATH
                                                    self.config["pdf_footer_image"] = FOOTER_IMAGE_PATH
                                                    self.pdf_footer_text = ""
                                                    self.config["pdf_footer_text"] = ""
                                                    self.save_configuration()
                                                    self._send_message(base_url, chat_id, "✅ PDF footer image updated.")
                                                else:
                                                    self._send_message(base_url, chat_id, "❌ Failed to download photo.")
                                            else:
                                                self._send_message(base_url, chat_id, "❌ Failed to get file info.")
                                        else:
                                            self._send_message(base_url, chat_id, "❌ Failed to fetch file info.")
                                    except Exception as e:
                                        self._send_message(base_url, chat_id, f"❌ Error: {e}")
                                else:
                                    if self.pdf_footer_image and os.path.exists(self.pdf_footer_image):
                                        try:
                                            os.remove(self.pdf_footer_image)
                                        except:
                                            pass
                                        self.pdf_footer_image = ""
                                        self.config["pdf_footer_image"] = ""
                                    new_footer = text[7:].strip() if len(text) > 7 else ""
                                    self.pdf_footer_text = new_footer
                                    self.config["pdf_footer_text"] = new_footer
                                    self.save_configuration()
                                    if new_footer:
                                        self._send_message(base_url, chat_id, "✅ PDF footer text updated.")
                                    else:
                                        self._send_message(base_url, chat_id, "✅ PDF footer text cleared.")
                                continue

                            # ---- /broadcast ----
                            if text.lower().startswith("/broadcast "):
                                broadcast_msg = text[11:].strip()
                                if broadcast_msg:
                                    self._broadcast_text(base_url, broadcast_msg, chat_id)
                                else:
                                    self._send_message(base_url, chat_id, "❌ Provide message.")
                                continue

                            # ---- Credits ----
                            if text.lower().startswith("/trecharge "):
                                amount_str = text[11:].strip()
                                try:
                                    amount = int(amount_str)
                                    if amount > 0:
                                        new_total = self.add_telegram_credits(amount)
                                        self._send_message(base_url, chat_id,
                                            f"✅ Recharged {amount} Telegram credits. Total: {new_total}")
                                        self.refresh_credits_gui()
                                    else:
                                        self._send_message(base_url, chat_id, "❌ Positive number required.")
                                except ValueError:
                                    self._send_message(base_url, chat_id, "❌ Provide a number.")
                                continue

                            if text.lower() == "/tbalance":
                                bal = self.get_telegram_credits()
                                self._send_message(base_url, chat_id, f"🤖 Telegram credits: {bal}")
                                continue

                            if text.lower().startswith("/recharge "):
                                amount_str = text[10:].strip()
                                try:
                                    amount = int(amount_str)
                                    if amount > 0:
                                        new_total = self.add_whatsapp_credits(amount)
                                        self._send_message(base_url, chat_id,
                                            f"✅ Recharged {amount} WhatsApp credits. Total: {new_total}")
                                        self.refresh_credits_gui()
                                    else:
                                        self._send_message(base_url, chat_id, "❌ Positive number required.")
                                except ValueError:
                                    self._send_message(base_url, chat_id, "❌ Provide a number.")
                                continue

                            if text.lower() == "/balance":
                                bal = self.get_whatsapp_credits()
                                self._send_message(base_url, chat_id, f"💰 WhatsApp credits: {bal}")
                                continue

                            # ---- /prompt ----
                            if text.lower() == "/prompt":
                                help_text = (
                                    "*Available Commands (Master only):*\n\n"
                                    "🔹 `/newbot <token>` – Change Telegram bot token.\n"
                                    "🔹 `/adduser <userid>` – Add admin user (auto‑send).\n"
                                    "🔹 `/remove <userid>` – Remove admin user.\n"
                                    "🔹 `/users` – Export Excel with admin_users and telegram_users.\n"
                                    "🔹 `/addgroup <group_id>` – Add a group for auto‑send.\n"
                                    "    *In a group, use `/addgroup` without argument to add this group.*\n"
                                    "🔹 `/removegroup <group_id>` – Remove a group.\n"
                                    "    *Use `/removegroup` without argument to list all added groups.*\n"
                                    "🔹 `/getconfig` – Send current config file.\n"
                                    "🔹 `/getdb` – Send database file.\n"
                                    "🔹 `/updateconfig` – Reply with a JSON file to replace config.\n"
                                    "🔹 `/message <text>` – Set caption message.\n"
                                    "🔹 `/footer <text>` – Set PDF footer text (or reply to photo).\n"
                                    "🔹 `/broadcast <message>` – Broadcast to ALL users.\n"
                                    "🔹 `/recharge <number>` – Add WhatsApp credits.\n"
                                    "🔹 `/balance` – Show WhatsApp credits.\n"
                                    "🔹 `/trecharge <number>` – Add Telegram credits.\n"
                                    "🔹 `/tbalance` – Show Telegram credits.\n"
                                    "🔹 `/prompt` – Show this help.\n\n"
                                    "📌 *For all users:*\n"
                                    "• `/start` – Get your ID and welcome message.\n"
                                    "• To request a report, send:\n"
                                    "  `[PATIENT ID]`\n"
                                    "  `[PATIENT FIRST NAME]`"
                                )
                                self._send_message(base_url, chat_id, help_text)
                                continue

                            # ---- Handle reply to /broadcast ----
                            if msg.get("reply_to_message") and text.lower() == "/broadcast":
                                reply_to = msg["reply_to_message"]
                                self._broadcast_reply(base_url, reply_to, chat_id)
                                continue

                            # ---- Handle file upload for /updateconfig ----
                            if expecting_config and msg.get("reply_to_message") and msg["reply_to_message"]["message_id"] == config_update_message_id:
                                document = msg.get("document")
                                if document and document.get("mime_type") == "application/json":
                                    file_id = document["file_id"]
                                    get_file_url = f"{base_url}/getFile?file_id={file_id}"
                                    file_resp = requests.get(get_file_url, timeout=10)
                                    if file_resp.status_code == 200:
                                        file_data = file_resp.json()
                                        if file_data.get("ok"):
                                            file_path = file_data["result"]["file_path"]
                                            download_url = f"https://api.telegram.org/file/bot{self.TELEGRAM_BOT_TOKEN}/{file_path}"
                                            content_resp = requests.get(download_url, timeout=30)
                                            if content_resp.status_code == 200:
                                                try:
                                                    new_config = content_resp.json()
                                                    required_keys = ["telegram_bot_token", "whatsapp_api_key", "institute_name"]
                                                    missing = [k for k in required_keys if k not in new_config]
                                                    if missing:
                                                        self._send_message(base_url, chat_id, f"❌ Invalid config: missing keys {missing}")
                                                    else:
                                                        backup_path = CONFIG_FILE + ".bak"
                                                        if os.path.exists(CONFIG_FILE):
                                                            shutil.copy2(CONFIG_FILE, backup_path)
                                                            self.log_message(f"📁 Config backed up to {backup_path}")
                                                        with open(CONFIG_FILE, "w") as f:
                                                            json.dump(new_config, f, indent=4)
                                                        self.load_configuration()
                                                        self.TELEGRAM_BOT_TOKEN = self.config["telegram_bot_token"]
                                                        self.update_bot_username()
                                                        self.refresh_bot_info_gui()
                                                        self.save_configuration()
                                                        self._send_message(base_url, chat_id, "✅ Config updated and reloaded successfully.")
                                                        self.log_message("✅ Config updated via Telegram.")
                                                except json.JSONDecodeError:
                                                    self._send_message(base_url, chat_id, "❌ Invalid JSON file.")
                                            else:
                                                self._send_message(base_url, chat_id, "❌ Failed to download file content.")
                                        else:
                                            self._send_message(base_url, chat_id, "❌ Failed to get file info.")
                                    else:
                                        self._send_message(base_url, chat_id, "❌ Failed to fetch file.")
                                else:
                                    self._send_message(base_url, chat_id, "❌ Please reply with a JSON file.")
                                expecting_config = False
                                continue

                        # ---------- Patient Query (any user, including groups) ----------
                        lines = [line.strip() for line in text.split("\n") if line.strip()]
                        if len(lines) >= 2:
                            query_id = lines[0]
                            query_name = lines[1].lower()

                            self.log_message(f"🔍 Searching pdf_index for patient ID: {query_id}, name prefix: {query_name[:4]}")
                            cached_pdf = self.get_saved_pdf_for_patient(query_id, query_name)
                            if cached_pdf:
                                cached_path, c_id, c_name, c_acc = cached_pdf
                                self.log_message(f"📄 PDF Found in cache: {os.path.basename(cached_path)}")
                                try:
                                    self._send_message(base_url, chat_id, "📄 Found existing report. Sending...")
                                    ok = self.dispatch_to_telegram(cached_path, c_id, c_name, c_acc, chat_id)
                                    if ok:
                                        self.root.after(0, lambda pi=c_id, pn=c_name, ac=c_acc, fp=cached_path:
                                                        self.upsert_grid_record(fp, pi, pn, ac, "📤 Sent via Bot (Cached)"))
                                        self._send_message(base_url, chat_id, "✅ Report sent successfully!")
                                        self.log_message(f"✅ Cached PDF sent to {chat_id}")
                                    else:
                                        self._send_message(base_url, chat_id, "❌ Failed to send PDF. Please try again later.")
                                        self.log_message(f"⚠️ Failed to send cached PDF for {c_id} - {c_name}")
                                except Exception as ex:
                                    self._send_message(base_url, chat_id, f"❌ Error sending cached report.")
                                    self.log_message(f"❌ Bot error (cached): {ex}")
                                continue

                            self.log_message(f"🔍 Cache miss. Searching dicom_index for patient ID: {query_id}, name prefix: {query_name[:4]}")
                            matched_entries = self.get_patient_files_from_db(query_id, query_name)

                            if matched_entries:
                                valid_entries = []
                                for entry in matched_entries:
                                    file_path, p_id, p_name, acc_no = entry
                                    if os.path.exists(file_path):
                                        valid_entries.append(entry)
                                    else:
                                        self._send_message(base_url, chat_id,
                                            f"⚠️ File for Patient `{p_name}` (ID: {p_id}) is missing.")
                                if not valid_entries:
                                    self._send_message(base_url, chat_id, "❌ Report Not Available")
                                    continue

                                self._send_message(base_url, chat_id,
                                    f"🔍 Found **{len(valid_entries)}** available DICOM file(s). Generating combined PDF...")
                                try:
                                    file_paths = [entry[0] for entry in valid_entries]
                                    _fp0, p_id, p_name, acc_no = valid_entries[0]
                                    study_date = ""
                                    try:
                                        first_ds = pydicom.dcmread(file_paths[0], stop_before_pixels=True)
                                        study_date = self._normalize_string(first_ds.get("StudyDate", ""))
                                    except Exception:
                                        pass
                                    bot_pdf_path = os.path.join(
                                        self.get_pdf_folder(),
                                        self.build_pdf_filename(p_name, study_date)
                                    )
                                    self.log_message(f"📄 Generating PDF from {len(file_paths)} DICOM files...")
                                    self.generate_pdf_report_from_dicom(file_paths, bot_pdf_path)
                                    ok = self.dispatch_to_telegram(bot_pdf_path, p_id, p_name, acc_no, chat_id)
                                    if ok:
                                        for entry in valid_entries:
                                            file_path, e_id, e_name, e_acc = entry
                                            self.root.after(0, lambda pi=e_id, pn=e_name, ac=e_acc, fp=file_path:
                                                            self.upsert_grid_record(fp, pi, pn, ac, "📤 Sent via Bot (Combined)"))
                                        self.save_pdf_index(p_id, p_name, acc_no, bot_pdf_path)
                                        self._send_message(base_url, chat_id,
                                            f"✅ Combined PDF with {len(valid_entries)} image(s) sent successfully!")
                                        self.log_message(f"✅ New PDF generated and sent to {chat_id}")
                                    else:
                                        self._send_message(base_url, chat_id, "❌ Failed to send PDF.")
                                except Exception as ex:
                                    self._send_message(base_url, chat_id, f"❌ Error generating combined PDF.")
                                    self.log_message(f"❌ Bot error: {ex}")
                            else:
                                self.log_message(f"❌ No record found for patient ID: {query_id}, name: {query_name}")
                                self._send_message(base_url, chat_id, "❌ Report Not Available")
            except Exception as e:
                self.log_message(f"Bot polling error: {e}")
            time.sleep(1)

    def _send_message(self, base_url, chat_id, text, parse_mode="Markdown"):
        try:
            requests.post(f"{base_url}/sendMessage", json={"chat_id": chat_id, "text": text, "parse_mode": parse_mode}, timeout=10)
        except Exception as e:
            self.log_message(f"Error sending message: {e}")

    def _broadcast_text(self, base_url, text, master_chat_id):
        users = self.get_all_telegram_users()
        if not users:
            self._send_message(base_url, master_chat_id, "No users to broadcast to.")
            return
        count = 0
        for uid in users:
            try:
                requests.post(f"{base_url}/sendMessage", json={"chat_id": uid, "text": text, "parse_mode": "Markdown"}, timeout=10)
                count += 1
            except Exception as e:
                self.log_message(f"Broadcast failed to {uid}: {e}")
        self._send_message(base_url, master_chat_id, f"✅ Broadcast sent to {count} users.")

    def _broadcast_reply(self, base_url, reply_msg, master_chat_id):
        users = self.get_all_telegram_users()
        if not users:
            self._send_message(base_url, master_chat_id, "No users to broadcast to.")
            return
        count = 0
        for uid in users:
            try:
                orig_chat_id = reply_msg["chat"]["id"]
                orig_msg_id = reply_msg["message_id"]
                copy_url = f"{base_url}/copyMessage"
                payload = {
                    "chat_id": uid,
                    "from_chat_id": orig_chat_id,
                    "message_id": orig_msg_id
                }
                if "caption" in reply_msg:
                    payload["caption"] = reply_msg["caption"]
                resp = requests.post(copy_url, json=payload, timeout=15)
                if resp.status_code == 200:
                    count += 1
                else:
                    self.log_message(f"Copy failed for {uid}: {resp.text}")
            except Exception as e:
                self.log_message(f"Broadcast reply error to {uid}: {e}")
            time.sleep(0.1)
        self._send_message(base_url, master_chat_id, f"✅ Broadcast reply sent to {count} users.")

    # ---------- Export Users to Excel ----------
    def _export_users_excel(self, base_url, master_chat_id):
        if not EXCEL_AVAILABLE:
            self._send_message(base_url, master_chat_id, "❌ openpyxl not installed. Please install 'openpyxl' for Excel export.")
            return

        try:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Admin Users"
            ws1.append(["User ID", "Added By", "Added At"])
            admin_users = self.get_admin_users()
            conn = sqlite3.connect(DATABASE_PATH)
            c = conn.cursor()
            for uid in admin_users:
                c.execute("SELECT added_by, added_at FROM admin_users WHERE user_id = ?", (uid,))
                row = c.fetchone()
                if row:
                    added_by = row[0]
                    added_at = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(row[1]))
                    ws1.append([uid, added_by, added_at])
            ws2 = wb.create_sheet("Telegram Users")
            ws2.append(["User ID", "Username", "Full Name", "First Seen"])
            c.execute("SELECT user_id, username, full_name, first_seen FROM telegram_users")
            for row in c.fetchall():
                first_seen = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(row[3])) if row[3] else ""
                ws2.append([row[0], row[1], row[2], first_seen])
            conn.close()

            temp_excel = os.path.join(os.environ.get('TEMP', '.'), f"users_{int(time.time())}.xlsx")
            wb.save(temp_excel)

            with open(temp_excel, "rb") as f:
                files = {"document": (os.path.basename(temp_excel), f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                payload = {"chat_id": master_chat_id, "caption": "📊 User list exported."}
                res = requests.post(f"{base_url}/sendDocument", data=payload, files=files, timeout=30)
                if res.status_code == 200:
                    self.log_message("✅ Users Excel sent.")
                else:
                    self.log_message(f"❌ Failed to send Excel: {res.text}")
            os.remove(temp_excel)
        except Exception as e:
            self.log_message(f"Error exporting users: {e}")
            self._send_message(base_url, master_chat_id, f"❌ Error exporting users: {e}")

if __name__ == "__main__":
    root = Tk()
    app = RadXrReceiverApp(root)
    root.mainloop()
